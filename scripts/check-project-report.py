import sys,json,re
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'.tools/report-python'))
sys.stdout.reconfigure(encoding='utf-8')
import pymupdf
from docx import Document
OUT=ROOT/'reports/project-documentation';doc=Document(OUT/'AI-Robot-Butler-Report.docx')
pdf=pymupdf.open(OUT/'AI-Robot-Butler-Report-browser.pdf')
norm=lambda s:re.sub(r'\s+','',s)
page_text=[norm(page.get_text()) for page in pdf]
body_start=None
for i,page in enumerate(pdf):
    if norm(page.get_text(sort=True)).startswith(norm('บทที่ 1 บทนำ')):
        body_start=i
    if body_start is not None:break
assert body_start is not None
targets=[p.text for p in doc.paragraphs if p.style.name.startswith('Heading') or p.style.name in ['Figure Caption','Table Caption']]
mapping={}
for target in targets:
    start=0 if target in ['หน้าอนุมัติ (รอข้อมูลและการลงนาม)','บทคัดย่อ','Abstract','กิตติกรรมประกาศ','สารบัญ','สารบัญภาพ','สารบัญตาราง'] else body_start
    for i in range(start,len(pdf)):
        if norm(target) in page_text[i]:mapping[target]=i+1;break
(OUT/'toc-pages.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2),encoding='utf-8')
missing=[target for target in targets if target not in mapping]
assert not missing,missing
full=''.join(page_text)
for row in json.loads((ROOT/'benchmarks/robot-ai/cases.json').read_text(encoding='utf-8'))['cases']:
    assert row['id'] in full,row['id']
assert 'คลิกขวาเพื่ออัปเดตสารบัญ' not in full
blank=[i+1 for i,page in enumerate(pdf) if len(norm(page.get_text()))<8]
assert not blank,blank
for name,index in [('preview-cover',0),('preview-toc',5),('preview-architecture',mapping['3.2 สถาปัตยกรรมของระบบ']-1),('preview-benchmark',mapping['5.3 ผลเปรียบเทียบการตีความคำสั่ง']-1),('preview-cases',mapping['ตารางที่ ข.2 ข้อความ ผลที่คาดหวัง และผล hybrid']-1)]:
    pdf[index].get_pixmap(matrix=pymupdf.Matrix(1.2,1.2)).save(OUT/f'{name}.png')
result={'pdfPages':len(pdf),'bodyStartsOnPdfPage':body_start+1,'wordTables':len(doc.tables),'wordImages':len(doc.inline_shapes),'mappedEntries':len(mapping),'benchmarkCases':40,'blankPages':blank}
(OUT/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result))

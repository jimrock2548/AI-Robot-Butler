"""Build the project report from repository evidence; no robot or model calls."""
import sys, json, hashlib
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.tools/report-python'))
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import pymupdf

OUT = ROOT / 'reports/project-documentation'
OUT.mkdir(parents=True, exist_ok=True)
TITLE = 'ระบบหุ่นยนต์ผู้ช่วยและนำทางด้วยปัญญาประดิษฐ์ผ่านเว็บแอปพลิเคชัน'
EN = 'Web-Based AI Robot Assistant and Navigation System'
reports = {name: json.loads((ROOT / f'reports/ai-benchmark/{name}.json').read_text(encoding='utf-8')) for name in ('baseline','hybrid','model')}
dataset = json.loads((ROOT / 'benchmarks/robot-ai/cases.json').read_text(encoding='utf-8'))
template_info = []
for path in (ROOT / 'ตัวอย่างรูปเล่ม').rglob('*'):
    if path.suffix.lower() == '.docx':
        sample = Document(path)
        template_info.append({'file': str(path.relative_to(ROOT)), 'type': 'docx', 'paragraphs': len(sample.paragraphs), 'sections': len(sample.sections), 'headings': [p.text for p in sample.paragraphs if p.text.startswith('บทที่')][:15]})
    elif path.suffix.lower() == '.pdf':
        with pymupdf.open(path) as sample:
            template_info.append({'file': str(path.relative_to(ROOT)), 'type': 'pdf', 'pages': len(sample), 'cover': sample[0].get_text()[:1200]})
(OUT / 'template-review.json').write_text(json.dumps(template_info, ensure_ascii=False, indent=2), encoding='utf-8')

# Original diagrams describing implemented components, not screenshots of a running robot.
font = ImageFont.truetype('C:/Windows/Fonts/tahoma.ttf', 23)
bold = ImageFont.truetype('C:/Windows/Fonts/tahomabd.ttf', 25)
def diagram(name, boxes, arrows, size=(1500,820)):
    im = Image.new('RGB', size, 'white'); d = ImageDraw.Draw(im)
    for a,b in arrows:
        d.line([a,b], fill='#4d6476', width=4)
        import math
        angle=math.atan2(b[1]-a[1],b[0]-a[0]); length=14
        d.polygon([b,(b[0]-length*math.cos(angle-.5),b[1]-length*math.sin(angle-.5)),(b[0]-length*math.cos(angle+.5),b[1]-length*math.sin(angle+.5))],fill='#4d6476')
    for rect,title,desc in boxes:
        d.rounded_rectangle(rect, radius=14, fill='#eef4f8', outline='#345a72', width=3)
        x1,y1,x2,y2=rect
        lines=[title]+desc.split('\n')
        for i,line in enumerate(lines):
            f=bold if i==0 else font; w=d.textbbox((0,0),line,font=f)[2]
            d.text(((x1+x2-w)/2,y1+18+i*34),line,font=f,fill='#16354a')
    im.save(OUT / f'{name}.png')
diagram('architecture', [((35,110,370,245),'Admin / Operator','Next.js :3000\nSession cookie'),((35,380,370,515),'Visitor Kiosk','Next.js :3001\nServer-side device token'),((560,250,935,440),'Backend','Express :4000\nPolicy + jobs + audit\nREST + Socket.IO'),((1120,35,1465,165),'PostgreSQL 16','Prisma / persistence'),((1120,225,1465,355),'Hermes Pro Max','SLAMTEC REST\nMap / pose / actions'),((1120,415,1465,545),'Local Qwen','llama-server :8080\nJSON intent / chat'),((1120,605,1465,735),'SearXNG (optional)','Internet search')], [((370,180),(560,305)),((370,445),(560,385)),((935,285),(1120,100)),((935,325),(1120,290)),((935,365),(1120,480)),((935,405),(1120,670))])
diagram('intent-flow',[((70,45,650,150),'Latest user message','Thai / English'),((70,220,650,325),'Rules + destination names','Stop / speech act / aliases'),((800,220,1410,325),'Model fallback','Structured JSON + allowed POI IDs'),((70,420,650,555),'Backend policy','Schema / ambiguity / rights\nFresh POI / map / robot state'),((800,420,1410,555),'Reject or clarify','No job / no motion'),((70,650,650,765),'Accepted request','dry-run preview OR validated job')],[((360,150),(360,220)),((650,272),(800,272)),((1110,325),(650,465)),((360,325),(360,420)),((650,485),(800,485)),((360,555),(360,650))])
diagram('data-model',[((45,65,465,205),'robots','id / code / base_url'),((540,65,965,205),'maps','robot_id / version / status'),((1050,65,1465,205),'pois + keepout_zones','map_id / geometry'),((45,325,465,465),'jobs','robot_id / destination / status'),((540,325,965,465),'job_events','job_id / at / detail'),((1050,325,1465,465),'users + sessions','roles / login sessions'),((45,600,465,745),'telemetry + control_lock','robot_id / measurements / owner'),((540,600,965,745),'audit_log','actor / action / detail'),((1050,600,1465,745),'AI logs','conversations / rejections')],[((465,130),(540,130)),((965,130),(1050,130)),((240,205),(240,325)),((465,395),(540,395)),((180,205),(180,600))])
im=Image.new('RGB',(1500,650),'white'); d=ImageDraw.Draw(im)
d.text((70,30),'Regression benchmark accuracy (40 cases)',font=bold,fill='#16354a')
for i,(name,label) in enumerate([('baseline','Previous policy'),('hybrid','Rules + model'),('model','Model + safety rules')]):
    score=reports[name]['summary']['accuracyPct']; y=130+i*150
    d.text((70,y),label,font=font,fill='#16354a'); d.rectangle((400,y,400+score*8,y+60),fill=['#9ba9b3','#2b7199','#477c60'][i]);d.text((420+score*8,y+15),f'{score:g}%',font=bold,fill='#16354a')
d.text((70,590),'PC / simulated robot / dry-run; not an unseen-language or physical-motion test.',font=font,fill='#4d6476'); im.save(OUT/'benchmark-accuracy.png')

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Inches(1.5);sec.left_margin=Inches(1.5);sec.right_margin=Inches(1);sec.bottom_margin=Inches(1)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
for name in ['Normal','Heading 1','Heading 2','Heading 3','Caption']:
    style=doc.styles[name];style.font.name='TH Sarabun New';style.font.size=Pt(16)
    rpr=style.element.get_or_add_rPr();fonts=OxmlElement('w:rFonts')
    for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'TH Sarabun New')
    rpr.append(fonts)
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.line_spacing=1.15
for name,size in [('Heading 1',20),('Heading 2',18),('Heading 3',16)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
doc.styles['Normal'].paragraph_format.first_line_indent=Cm(1.25)
for name in ['Figure Caption', 'Table Caption']:
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles['Caption']
doc.core_properties.title=TITLE;doc.core_properties.subject='รายงานโครงงานฉบับร่างตามหลักฐานใน repository ณ 12 ตุลาคม 2569';doc.core_properties.author='[กรอกชื่อผู้จัดทำ]'
text_source=[]; figures=[]; tables=[]
def p(text,center=False):
    text_source.append(text);para=doc.add_paragraph(text)
    if center:para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.first_line_indent=Pt(0)
    return para
def h(text,level=2):
    text_source.append('#'*level+' '+text);doc.add_heading(text,level=level)
def page(title):
    doc.add_page_break();h(title,1)
def table(caption,headers,rows):
    tables.append(caption);p(caption).style='Table Caption';t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid'
    for cell,txt in zip(t.rows[0].cells,headers):cell.text=str(txt)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell,txt in zip(t.add_row().cells,row):cell.text=str(txt)
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.first_line_indent=Pt(0);para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=1
                for run in para.runs:run.font.size=Pt(14)
    text_source.append('\n'.join(' | '.join(map(str,row)) for row in [headers,*rows]));p('')
def figure(name,caption):
    figures.append(caption);doc.add_picture(str(OUT/f'{name}.png'),width=Cm(14.5));p(caption,True).style='Figure Caption'
def field(para,code):
    run=para.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');run._r.append(begin)
    instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=code;run._r.append(instr)
    separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate');run._r.append(separate)
    run=para.add_run('คลิกขวาเพื่ออัปเดตสารบัญ');end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');run._r.append(end)

p('รายงานโครงงานฉบับร่าง',True);p('');p(TITLE,True);p(EN,True);p('');p('');p('[กรอกชื่อ–นามสกุลผู้จัดทำ และรหัสนักศึกษา]',True);p('');p('');p('ปริญญานิพนธ์นี้เป็นส่วนหนึ่งของการศึกษาตามหลักสูตร',True);p('[กรอกชื่อปริญญาและสาขาวิชา]',True);p('[กรอกภาควิชา คณะ และมหาวิทยาลัย]',True);p('ปีการศึกษา [กรอกปีการศึกษา]',True);p('ข้อมูลโครงการ ณ วันที่ 12 ตุลาคม 2569',True)
page('หน้าอนุมัติ (รอข้อมูลและการลงนาม)')
for text in [f'ชื่อโครงงาน: {TITLE}',f'ชื่อภาษาอังกฤษ: {EN}','ผู้จัดทำ: [กรอกชื่อและรหัสนักศึกษา]','สาขาวิชา / ภาควิชา: [กรอกข้อมูล]','คณะ / มหาวิทยาลัย: [กรอกข้อมูล]','อาจารย์ที่ปรึกษา: [กรอกชื่อและตำแหน่ง]','ปีการศึกษา: [กรอกปีการศึกษา]'] :p(text)
p('หน้านี้เป็นแบบร่างสำหรับกรอกข้อมูล ไม่ใช่หลักฐานการอนุมัติหรือรับรองผลโครงงาน')
for role in ['ประธานกรรมการ','กรรมการ','อาจารย์ที่ปรึกษา','หัวหน้าภาควิชา']:p(f'........................................................ {role}\n([กรอกชื่อและตำแหน่ง])')
page('บทคัดย่อ')
p('โครงงานนี้พัฒนาระบบหุ่นยนต์ผู้ช่วยและนำทางผ่านเว็บแอปพลิเคชัน โดยแยกเว็บไซต์สำหรับเจ้าหน้าที่ออกจากเว็บไซต์สำหรับผู้มาเยือน และใช้ Backend กลางประสานการควบคุมหุ่นยนต์ Hermes Pro Max ผ่าน REST API ระบบประกอบด้วยการแสดงแผนที่ SLAM และตำแหน่งแบบต่อเนื่อง การจัดการจุดสนใจ การควบคุมสิทธิ์และงานนำทาง ตลอดจนการรับคำขอภาษาธรรมชาติด้วยโมเดล Qwen3.5-2B ที่ทำงานบนเครื่องพัฒนา')
p('การตีความคำสั่งใช้กฎภาษาและชื่อเรียกอื่นของจุดสนใจร่วมกับโมเดลที่คืนผลแบบ JSON โดย Backend ตรวจเจตนา รหัสจุดหมาย สิทธิ์ สถานะหุ่น และข้อมูลแผนที่ก่อนสร้างงาน ระบบไม่อนุญาตให้โมเดลกำหนดพิกัดอิสระ กรณีชื่อกำกวมหรือไม่พบจุดหมายจะขอให้ระบุใหม่ และคำหยุดที่รองรับไม่ต้องรอการประมวลผลของโมเดล')
p('ผลทดสอบย้อนหลังบนเครื่องพัฒนาในวันที่ 7 ตุลาคม 2569 พบว่าชุด regression จำนวน 40 กรณีผ่านเพิ่มจาก 21 กรณี หรือร้อยละ 52.5 เป็น 40 กรณี หรือร้อยละ 100 สำหรับวิธีผสมกฎกับโมเดล ไม่พบการเสนอเดินผิดเจตนาหรือผิดจุดหมายในชุดนี้ การทดสอบเป็น dry-run กับหุ่นจำลอง และไม่ใช่การประเมินความสามารถกับประโยคใหม่ที่ไม่เคยใช้พัฒนา มีผลตรวจ Backend ผ่าน 51 รายการและผล typecheck ผ่านตามบันทึกโครงการ')
p('ขณะจัดทำรายงานไม่ได้เชื่อมต่อหุ่นจริง ผลอ่านแผนที่และ POI จากฮาร์ดแวร์เป็นหลักฐานจากรอบก่อนหน้า ยังไม่มีผลตรวจรับการเคลื่อนที่จริง การประมวลผลบน Jetson และระบบเสียงยังไม่เสร็จ รวมถึงบริบทถามตอบต่อเนื่องสำหรับคำสั่งนำทางและโหมดสำรวจจริงที่ยังเป็นงานถัดไป')
p('คำสำคัญ: หุ่นยนต์ผู้ช่วย, การนำทาง, SLAM, จุดสนใจ, ภาษาธรรมชาติ, เว็บแอปพลิเคชัน')
page('Abstract')
p('This project develops a web-based robot assistant and navigation system integrating a Hermes Pro Max robot through a central backend. Separate administrator and visitor websites support access control, live SLAM map visualization, point-of-interest management, navigation jobs, and natural-language requests. Qwen3.5-2B runs locally on the development PC.')
p('The command pipeline combines language rules and destination aliases with structured model output. The backend validates intent, known destination identifiers, permissions, fresh map data, and robot readiness before creating a job. Ambiguous requests require clarification. Supported stop requests bypass model inference.')
p('Recorded regression results from October 7, 2026 show an increase from 21 of 40 passing cases (52.5%) to 40 of 40 (100%) using the hybrid strategy, with no incorrect proposed movement in this dataset. All benchmark requests used a simulated robot and dry-run mode. These results do not establish generalization to unseen language or successful physical navigation. The project record also reports 51 passing backend tests.')
p('The robot was disconnected while preparing this report. Physical-motion acceptance, deployment and measurement on Jetson, speech input/output, conversational navigation context, and actual exploration-mode control remain incomplete. The report distinguishes implementation evidence from planned capabilities.')
p('Keywords: robot assistant, navigation, SLAM, points of interest, natural language, web application')
page('กิตติกรรมประกาศ')
p('[กรอกข้อความขอบคุณบุคคล หน่วยงาน และผู้สนับสนุนที่มีส่วนร่วมจริง พร้อมตรวจสอบชื่อและตำแหน่งก่อนส่งเล่ม]')
p('หมายเหตุผู้จัดทำเอกสาร: ไม่ได้แต่งชื่อผู้สนับสนุน การอนุมัติ หรือผลทดลองที่ยังไม่มีหลักฐาน เอกสารนี้เป็นร่างที่จัดทำจากโค้ด แผนงาน ผล benchmark และรูปแบบตัวอย่างที่ผู้ใช้ให้ไว้')
page('สารบัญ');field(doc.add_paragraph(),' TOC \\o "1-2" \\h \\z \\u ')
page('สารบัญภาพ');field(doc.add_paragraph(),' TOC \\t "Figure Caption,1" \\h \\z ')
page('สารบัญตาราง');field(doc.add_paragraph(),' TOC \\t "Table Caption,1" \\h \\z ')

body=doc.add_section(WD_SECTION_START.NEW_PAGE);body.footer.is_linked_to_previous=False;body.header.is_linked_to_previous=False
num=OxmlElement('w:pgNumType');num.set(qn('w:start'),'1');body._sectPr.append(num)
footer=body.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(footer,' PAGE ')
h('บทที่ 1 บทนำ',1)
h('1.1 ที่มาและความสำคัญของปัญหา')
p('การให้ข้อมูลและนำทางผู้มาเยือนภายในอาคารต้องอาศัยข้อมูลสถานที่ที่ถูกต้องและการประสานงานกับเจ้าหน้าที่ ระบบหุ่นยนต์สามารถเป็นช่องทางให้ผู้ใช้เลือกจุดหมายและรับข้อมูลสถานะการเดินทางได้ แต่การต่อเว็บเข้ากับหุ่นโดยตรงทำให้ยากต่อการควบคุมสิทธิ์ การจัดการคำสั่งซ้อน และการตรวจสอบย้อนหลัง โครงงานจึงออกแบบ Backend เป็นตัวกลางที่ตรวจคำขอทุกครั้งก่อนส่งไปยังอุปกรณ์')
p('ผู้ใช้ควรสั่งงานได้ด้วยคำขอที่คุ้นเคย เช่น “ช่วยพาไปที่หน้าประตูหน่อย” แทนการจำรูปประโยคเดียว อย่างไรก็ตาม ความคลาดเคลื่อนของโมเดลภาษาอาจนำไปสู่การเลือกจุดหมายผิด จึงต้องแยกการตีความภาษาจากการอนุมัติคำสั่งเคลื่อนที่ โดยยึดรายการ POI และสถานะจากหุ่นเป็นข้อมูลอ้างอิง')
h('1.2 วัตถุประสงค์')
for text in ['พัฒนาเว็บไซต์ผู้ดูแลและเว็บไซต์ผู้ช่วยที่ใช้ Backend เดียวกัน พร้อมบทบาทและสิทธิ์ที่ตรวจสอบได้','แสดงแผนที่ ตำแหน่ง LiDAR และจุดสนใจจากหุ่น พร้อมจัดการ POI และงานนำทาง','รองรับคำขอไทยและอังกฤษหลายสำนวนด้วยกฎ ชื่อเรียกอื่น และโมเดล local ที่คืนผลแบบ structured output','สร้างชุดทดสอบและ benchmark ที่ตรวจการเลือกคำสั่งและจุดหมายโดยไม่ส่งคำสั่งเดินจริง','จัดทำหลักฐานและแผนตรวจรับเพื่อแยกความสำเร็จด้านซอฟต์แวร์ออกจากพฤติกรรมฮาร์ดแวร์']:p(text)
h('1.3 ขอบเขตของโครงงาน')
p('ขอบเขตปัจจุบันรองรับหุ่น Hermes Pro Max หนึ่งตัว ระบบผู้ใช้ admin/operator/viewer และจอ Kiosk ที่ใช้ device token มีหุ่นจำลองสำหรับพัฒนาและฐานข้อมูลแบบ memory หรือ PostgreSQL ส่วน AI ใช้ Qwen3.5-2B Q5_K_M ผ่าน llama-server บนพีซี การใช้ Jetson Xavier NX 8GB เป็นเป้าหมายการย้ายระบบในระยะถัดไป ไม่ใช่สภาพแวดล้อมของผล benchmark ที่รายงาน')
p('เว็บมีตัวเลือกโหมดปกติและโหมดสำรวจที่บันทึกในเบราว์เซอร์ แต่ยังไม่เปลี่ยนสถานะการสร้างแผนที่บนหุ่น ไม่มีการเริ่มสำรวจเมื่อเลือกโหมดหรือเมื่อเชื่อมต่อใหม่ งานที่ยังไม่อยู่ในขอบเขตที่เสร็จแล้ว ได้แก่ การสำรวจอัตโนมัติ ระบบเสียง STT/TTS การจดจำคำถามนำทางที่ค้าง และคำขอหลายขั้นตอน')
h('1.4 ประโยชน์ที่คาดว่าจะได้รับ')
p('ผู้ควบคุมมองเห็นข้อมูลแผนที่และสถานะจากเว็บเดียว ผู้มาเยือนมีเว็บไซต์เฉพาะสำหรับเลือกจุดหมายและพิมพ์คำขอ นักพัฒนาสามารถตรวจความถูกต้องของการตีความก่อนทดลองเคลื่อนที่ และมีบันทึกเหตุผลเมื่อระบบปฏิเสธคำขอ ผลที่คาดว่าจะได้รับในหน้างานต้องยืนยันด้วยการตรวจรับจริงก่อนสรุปประสิทธิผลด้านการบริการ')
h('1.5 นิยามศัพท์และการอ่านสถานะ')
table('ตารางที่ 1.1 นิยามศัพท์สำคัญ',['คำศัพท์','ความหมายในโครงงาน'],[['POI','จุดสนใจที่มีรหัส ชื่อ พิกัดและทิศบนแผนที่'],['SLAM','การสร้างแผนที่และประมาณตำแหน่งภายในสภาพแวดล้อม'],['Localization','การประมาณตำแหน่งหุ่นเทียบกับแผนที่ที่มีอยู่'],['dry-run','ตีความและแสดงผลโดยไม่สร้างงานหรือส่งคำสั่งเคลื่อนที่'],['ทำแล้ว','มีเส้นทางในโค้ดและหลักฐานทดสอบตามขอบเขตที่ระบุ'],['รอตรวจรับ','มีโค้ดรองรับแต่ยังไม่ยืนยันพฤติกรรมจริงของอุปกรณ์']])

page('บทที่ 2 แนวคิดและเทคโนโลยีที่เกี่ยวข้อง')
h('2.1 หุ่นยนต์เคลื่อนที่ แผนที่ และการระบุตำแหน่ง')
p('แผนที่สำหรับหุ่นต่างจากภาพแปลนอาคารที่ใช้ตกแต่งหน้าเว็บ แผนที่กริดมีจุดกำเนิด ความละเอียด และค่าแต่ละเซลล์สำหรับอ้างพิกัดจริง เมื่อแสดงบนจอจำเป็นต้องแปลงพิกัดและแกนให้ตรงกัน รวมทั้งรักษาจุดกำเนิดที่เป็นค่าติดลบ โครงงานอ่านข้อมูลกริดและตำแหน่งจากบริการ SLAM ของหุ่นตาม API ของผู้ผลิต [1] โดยไม่ได้พัฒนาอัลกอริทึม SLAM ใหม่')
p('การสร้างแผนที่กับการนำทางเป็นคนละสถานะ โค้ดนำทางกำหนดให้ mappingEnabled เป็น false และคุณภาพตำแหน่งผ่านเกณฑ์ก่อนส่งงาน การแสดง LiDAR ช่วยให้ผู้ควบคุมเห็นข้อมูลการสแกน แต่ไม่ใช่หลักฐานว่าระบบเว็บมีอัลกอริทึมหลบสิ่งกีดขวางเอง การวางแผนเส้นทางและการควบคุมฐานเคลื่อนที่ยังเป็นหน้าที่ของหุ่น')
h('2.2 เว็บแอปพลิเคชันและการแสดงข้อมูลต่อเนื่อง')
p('ระบบใช้ Next.js และ React สำหรับ Admin กับ Kiosk และใช้ Express เป็น Backend การแยกสองเว็บไซต์ช่วยให้จอผู้มาเยือนมีขั้นตอนเฉพาะของตนเอง ส่วนเจ้าหน้าที่เห็นการควบคุม แผนที่ งาน และบันทึก ระบบ REST รองรับคำขอที่ต้องมีผลยืนยัน และ Socket.IO ส่งข้อมูลสถานะจาก Backend ไปยังเว็บ')
p('ข้อมูลที่ส่งต่อเนื่องไม่ได้หมายความว่าจะรับได้ทุกข้อความ Socket.IO อธิบายการเรียงลำดับข้อความและข้อจำกัดการส่งแบบ at-most-once ตามค่าเริ่มต้น [3] โครงงานจึงแสดงเวลาอัปเดตและมี HTTP fallback แทนการตีความว่าการเชื่อมต่อเครือข่ายที่เปิดอยู่ทำให้ข้อมูลสดเสมอ')
h('2.3 โมเดลภาษา local และ structured output')
p('Qwen3.5-2B เป็นโมเดลที่มี model card เผยแพร่โดย Qwen [2] ในโครงงานใช้ไฟล์ GGUF แบบ Q5_K_M ผ่าน llama.cpp โดยจำกัดผลลัพธ์คำสั่งเป็น JSON แทนข้อความอิสระ ค่า confidence เป็นผลที่โมเดลรายงาน ไม่ใช่ความน่าจะเป็นที่ผ่านการปรับเทียบหรือการรับประกันว่าคำสั่งถูกต้อง')
p('การใช้ schema ช่วยกำหนดชนิดข้อมูลและชุดค่า intent ที่ยอมรับได้ แต่ไม่ได้รับรองความหมายของประโยค โค้ดจึงตรวจผลอีกชั้นด้วย Zod และตรวจว่ารหัส POI อยู่ในรายการที่อนุญาต การเสนอคำสั่งด้วยโมเดลไม่เท่ากับการอนุญาตให้ส่งคำสั่งหุ่นโดยตรง')
h('2.4 การจับคู่ชื่อและการแยกเจตนาภาษา')
p('วิธีแรกใช้รูปประโยคคำขอและชื่อเรียกอื่นที่จัดการไว้ใน metadata.aliases วิธีที่สองใช้โมเดลสำหรับประโยคที่กฎยังไม่รองรับ ทั้งสองวิธีแยก speechAct เป็น request, question, negated, hypothetical, quoted หรือ other เพื่อแยก “ช่วยพาไป” ออกจาก “อยู่ที่ไหน” และ “อย่าเพิ่งไป”')
p('การจับคู่ชื่อให้ความสำคัญกับชื่อเฉพาะที่ยาวกว่าเมื่อเกิดทับซ้อนในตำแหน่งเดียวกัน เช่น ห้องประชุม A เทียบกับห้องประชุม แต่ถ้าผู้ใช้ระบุสองจุดต่างตำแหน่งในประโยคต้องยังถือว่ามีหลายจุดหมาย ชื่อภาษาอังกฤษต้องตรวจขอบเขตคำ และชื่อไทยต้องไม่เปลี่ยนรูปสระจนทำให้คำเดิมจับไม่ตรง')
h('2.5 ฐานข้อมูล สิทธิ์ และการตรวจสอบย้อนหลัง')
p('PostgreSQL เก็บข้อมูลผ่าน Prisma โดยมีผู้ใช้ เซสชัน แผนที่ POI งาน เหตุการณ์ telemetry และบันทึกการทำงาน Prisma schema ใน repository เป็นหลักฐานโครงสร้างจริง ไม่ได้สรุปว่า ORM ทำให้ทุกคำสั่งปลอดภัยโดยอัตโนมัติ การตรวจบทบาทและสถานะงานอยู่ใน middleware และ service ของระบบ')
p('สิทธิ์ของ Admin ใช้เซสชันผู้ใช้ ส่วน Kiosk ใช้ token ที่ proxy เก็บฝั่งเซิร์ฟเวอร์ การบันทึก audit และเหตุผลปฏิเสธช่วยทบทวนว่าใครสั่งงานและทำไมคำขอไม่ผ่าน ระบบปัจจุบันบันทึก intent และ rejection แต่ยังไม่ได้เก็บแชตทั่วไปทุกข้อความเป็นประวัติถาวร')

page('บทที่ 3 การวิเคราะห์และออกแบบระบบ')
h('3.1 ผู้ใช้งานและความต้องการของระบบ')
table('ตารางที่ 3.1 ผู้ใช้งานและความสามารถหลัก',['ผู้ใช้','ความสามารถ','ข้อจำกัด'],[['ผู้มาเยือนผ่าน Kiosk','เลือก POI สาธารณะ พิมพ์ถาม/นำทาง หยุดงานของ Kiosk','ไม่เห็นจุดส่วนตัวและหยุดงานเจ้าหน้าที่ไม่ได้'],['viewer','ดูสถานะและข้อมูลตามสิทธิ์','ไม่สั่งงานควบคุม'],['operator','ควบคุม นำทาง และจัดการ POI','ต้องผ่าน lock และข้อกำหนดหุ่น'],['admin','จัดการผู้ใช้และควบคุมตามสิทธิ์ผู้ดูแล','คำสั่งยังต้องผ่านข้อกำหนดการนำทาง']])
p('ความต้องการที่สำคัญคือการมีงานนำทางที่กำลังทำเพียงงานเดียว การถือสิทธิ์ควบคุมด้วยมือ การหยุดที่ไม่รอ AI และการแยกข้อมูลจำลองจากข้อมูลหุ่นจริง ระบบต้องแจ้งเมื่อไม่สามารถอ่านข้อมูลได้แทนการเติมค่าที่ดูเหมือนจริง เช่น ความเร็วที่ firmware ไม่มี endpoint ให้ต้องแสดงว่าไม่ทราบ')
h('3.2 สถาปัตยกรรมของระบบ')
figure('architecture','ภาพที่ 3.1 สถาปัตยกรรมระบบตามโค้ดใน repository')
p('Admin และ Kiosk เรียก Backend เท่านั้น ไม่เรียกหุ่นจากเบราว์เซอร์โดยตรง Backend เลือก adapter ระหว่างตัวจำลองกับ Hermes และประสานฐานข้อมูล โมเดล และบริการค้นหา การเชื่อมต่อ SearXNG เป็นส่วนเสริมที่ต้องเปิดบริการและมีอินเทอร์เน็ต ไม่จำเป็นต่อ benchmark คำสั่งหุ่น')
h('3.3 การไหลของคำสั่งภาษาธรรมชาติ')
figure('intent-flow','ภาพที่ 3.2 การตีความ ตรวจคำขอ และสร้างงานนำทาง')
p('เมื่อรับข้อความ ระบบตรวจคำหยุดก่อนเพื่อหลีกเลี่ยงการรอโมเดล จากนั้นอ่าน POI ที่มีพิกัดถูกต้องและกรองจุดส่วนตัวตามช่องทาง กฎภาษาเลือกคำสั่งที่ชัดเจนได้ทันที ส่วนข้อความที่ยังไม่ชัดเจนส่งให้โมเดลพร้อมรายการปลายทาง โมเดลต้องคืน intent, confidence, poiCode, speechAct และ needsClarification')
p('Backend ปฏิเสธคำปฏิเสธ สมมติ และข้อความยกมา ตรวจชื่อจุดหมายและความกำกวม และอ่าน POI ใหม่ก่อน dispatch เพื่อไม่ใช้พิกัดที่ถูกแก้ระหว่างรอโมเดล หากเป็น dry-run จะคืนเพียงผลตีความ งานจริงต้องผ่านการตรวจแผนที่ ความพร้อมหุ่น และสิทธิ์เพิ่มเติม')
h('3.4 การออกแบบฐานข้อมูล')
figure('data-model','ภาพที่ 3.3 กลุ่มข้อมูลและความสัมพันธ์หลักตาม Prisma schema')
p('ภาพนี้เป็นภาพรวมความสัมพันธ์ที่ประกาศใน schema ไม่ใช่ ER diagram ที่แสดง foreign key ทุกฟิลด์ รหัส POI จากหุ่นจริงเป็น UUID ภายนอก ส่วน POI ในฐานข้อมูลของระบบมี id ภายในและ code จึงต้องไม่ใช้หมายเลขแถวแทนรหัสจุดหมายของหุ่น')
table('ตารางที่ 3.2 พจนานุกรมข้อมูลหลัก',['ตาราง','คีย์/ข้อมูลสำคัญ','หน้าที่'],[['users','id, username, password_hash, role, is_active','บัญชีและสิทธิ์'],['sessions','sid, sess, expire','เซสชันเข้าสู่ระบบ'],['robots','id, code, base_url','ข้อมูลการเชื่อมต่อหุ่น'],['maps','robot_id, version, status, geometry','แผนที่และรุ่นที่เผยแพร่'],['pois','map_id, code, name_th, x, y, heading_deg','จุดหมายของระบบ'],['keepout_zones','map_id, polygon','พื้นที่ตามข้อมูลแผนที่'],['jobs','robot_id, dest_poi_code, source, status','งานนำทาง'],['job_events','job_id, at, kind, detail','ลำดับเหตุการณ์งาน'],['telemetry_samples','robot_id, at, battery_pct, x, y','ข้อมูลสถานะตามเวลา'],['control_lock','robot_id และผู้ถือสิทธิ์','การควบคุมร่วมกัน'],['audit_log','ผู้กระทำ เหตุการณ์ เวลา','ตรวจสอบย้อนหลัง'],['ai_conversations / ai_rejections','transcript, intent, latency / reason','บันทึกการตีความและปฏิเสธ']])
h('3.5 ข้อกำหนดก่อนส่งคำสั่งเคลื่อนที่')
table('ตารางที่ 3.3 เงื่อนไขสำคัญของคำสั่งนำทาง',['รายการตรวจ','เงื่อนไขในโค้ด','เมื่อไม่ผ่าน'],[['อายุแผนที่','ไม่เกิน 5 วินาที','ปฏิเสธข้อมูลเก่า'],['โหมดแผนที่','mappingEnabled ต้องเป็น false','แจ้งให้กลับโหมดระบุตำแหน่ง'],['คุณภาพตำแหน่ง','อย่างน้อย 50 ตามค่าที่ตั้งไว้','ไม่สร้างงานนำทาง'],['จุดหมาย','พิกัด finite และอยู่ในขอบแผนที่','ปฏิเสธจุดนอกแผนที่'],['เซลล์พื้นที่ว่าง','สำหรับจุดทั่วไป ต้อง 0 < cell < 128','ปฏิเสธพื้นที่ไม่ทราบ/ไม่ว่าง'],['งานและ manual lock','ไม่มีความขัดกันกับงานหรือสิทธิ์','แจ้งหุ่นไม่พร้อม'],['สิทธิ์ Kiosk','POI ไม่เป็น is_public=false','ไม่ให้เลือกจุดส่วนตัว']])
p('ค่าจำกัดความเร็วที่เคยอ่านกลับในห้องแลปคือ 0.1 เมตรต่อวินาทีและ 0.3 เรเดียนต่อวินาที โค้ดตรวจการยืนยันและอ่านกลับก่อนถือว่าใช้งานได้ การเดินจริงยังต้องตรวจในพื้นที่ภายใต้การควบคุมของผู้ใช้ โครงการไม่ได้อ้างว่าชุดทดสอบซอฟต์แวร์แทนการตรวจรับทางกายภาพได้')
h('3.6 การเปลี่ยนสถานะงานและการยกเลิก')
p('Prisma schema กำหนดสถานะ queued, navigating, arrived, cancelled, failed และ stopped เมื่อมีคำหยุดระหว่าง inference ระบบเพิ่ม command epoch เพื่อทำให้ผลที่ตอบกลับช้ากลายเป็นคำขอที่ถูกยกเลิกและไม่ dispatch ภายหลัง การปล่อย manual lock ระหว่างตรวจความเร็วก็ต้องป้องกันคำสั่งเดินที่มาช้าเช่นเดียวกัน')
h('3.7 การออกแบบโหมดปกติและโหมดสำรวจ')
p('ตัวเลือกโหมดในหน้า Control ใช้ localStorage ของเบราว์เซอร์ เป็น preference ฝั่งเว็บ ไม่ใช่สถานะของหุ่นที่ได้รับการยืนยัน การโหลดหน้าและเลือกโหมดไม่มีการส่ง API หุ่นจริง การพัฒนาระยะถัดไปควรแยก requested mode และ applied mode มีการอ่านกลับยืนยัน และไม่เปลี่ยนโหมดโดยอัตโนมัติหลังเครือข่ายกลับมา')

page('บทที่ 4 การพัฒนาระบบ')
h('4.1 เครื่องมือและสภาพแวดล้อม')
table('ตารางที่ 4.1 เครื่องมือและอุปกรณ์ที่ใช้',['ส่วนงาน','เครื่องมือ/อุปกรณ์','สถานะ'],[['เว็บ','Next.js 14, React, TypeScript','Admin และ Kiosk แยกเว็บไซต์'],['Backend','Express, Socket.IO, Zod','REST และข้อมูลสถานะ'],['ฐานข้อมูล','PostgreSQL 16, Prisma 5.22','มี memory mode สำหรับทดสอบ'],['โมเดล local','Qwen3.5-2B Q5_K_M, llama.cpp b10545','ผล benchmark บน PC'],['หุ่น','Hermes Pro Max, SLAMTEC REST','หลักฐานอ่านค่าจากรอบก่อน'],['อุปกรณ์เป้าหมาย AI','Jetson Xavier NX 8GB','เตรียมแล้ว แต่ build พัก'],['การทดสอบ','Node test runner, TypeScript, benchmark script','หุ่นจำลองและ dry-run']])
p('ผล benchmark ระบุเครื่อง Windows x64, Intel Core i7-10700F และ RAM รวมประมาณ 63.9 GiB ค่า RAM นี้เป็นขนาดหน่วยความจำของเครื่อง ไม่ใช่ค่าการใช้หน่วยความจำสูงสุดของโมเดล เอกสารนี้ไม่อนุมานว่าความหน่วงบนพีซีจะเท่ากับ Jetson')
h('4.2 โครงสร้างโปรเจ็กต์และการแยกความรับผิดชอบ')
table('ตารางที่ 4.2 ตำแหน่งส่วนงานใน repository',['ตำแหน่ง','ความรับผิดชอบ'],[['apps/admin','หน้าเจ้าหน้าที่ แผนที่ ควบคุม งาน ผู้ใช้ และบันทึก AI'],['apps/kiosk','หน้าเลือกจุดหมายและแชตผู้มาเยือน'],['apps/backend/src/routes','REST endpoints และ middleware สิทธิ์'],['apps/backend/src/services','งาน หุ่น AI แผนที่ และข้อกำหนดนำทาง'],['apps/backend/src/hermes','adapter REST และการแปลงแผนที่'],['apps/backend/src/sim','หุ่นจำลอง'],['packages/shared','types และ schema ร่วม'],['benchmarks/robot-ai','ชุดข้อมูลทดสอบภาษา'],['reports/ai-benchmark','ผลรายกรณี JSON และ Markdown']])
h('4.3 แผนที่สดและการจัดการ POI')
p('LiveRobotMap แสดงกริด SLAM พร้อมหุ่น แท่นชาร์จ จุดสนใจ และ LiDAR ผู้ใช้ซูมและลากแผนที่ได้ Backend ส่ง snapshot ประมาณทุก 1 วินาทีและ telemetry ประมาณ 0.5 วินาที มี HTTP fallback ประมาณทุก 2 วินาที ช่วงเวลาเหล่านี้เป็นค่าการตั้งในระบบ ไม่ใช่หลักฐานว่าความหน่วงปลายทางจะคงที่ในทุกเครือข่าย')
p('การเพิ่มหรือแก้ POI ตรวจสิทธิ์ โหมดแผนที่ พิกัด และงาน/lock ก่อนเขียน จากนั้นอ่านรายการบนหุ่นกลับเพื่อตรวจชื่อ พิกัด ทิศและ aliases การลบตรวจเฉพาะ id ที่ระบุและอ่านยืนยันว่าหายไป ข้อมูล metadata เดิมถูกเก็บไว้เมื่อแก้ไข จุดสนใจบนหุ่นและแผนที่ที่เผยแพร่ในฐานข้อมูลเป็นแหล่งข้อมูลที่มีบทบาทต่างกัน')
h('4.4 การควบคุมด้วยมือและกลับจุดชาร์จ')
p('ผู้ใช้ที่มีสิทธิ์ต้องถือ manual lock จึงส่งคำสั่ง joystick ได้ lock มี TTL 30 วินาทีและ heartbeat ประมาณทุก 10 วินาที การปล่อยสิทธิ์และการหยุดส่งคำสั่งต้องทำให้การควบคุมสิ้นสุดตามเส้นทางที่ทดสอบ การกลับจุดชาร์จใช้ตำแหน่ง home ที่อ่านจากหุ่นและมีเงื่อนไขความพร้อมเช่นเดียวกับการนำทาง โหมดกลับแท่นอัตโนมัติปิดไว้ตามการใช้งานในห้องแลป')
h('4.5 การพัฒนา AI ภาษาธรรมชาติ')
p('robotLanguage.ts ทำหน้าที่ normalize แบบ NFC ตรวจรูปประโยค จับชื่อ/aliases และแยก speech act ส่วน robotIntent.ts ใช้กฎที่ชัดเจนก่อน แล้วเรียกโมเดลในกรณีอื่น prompt บอกรายการปลายทางและห้ามสร้างรหัสหรือพิกัดใหม่ schema จำกัด intent และ poiCode ขณะที่ Zod ตรวจผลตอบกลับอีกครั้ง')
table('ตารางที่ 4.3 โครงสร้างผลตีความคำสั่ง',['ฟิลด์','ชนิดข้อมูล','การใช้งาน'],[['intent','ค่าจากชุดที่กำหนด','navigate_to, where_is, return_to_dock, cancel_job, robot_status, answer_question, unknown'],['confidence','number 0–1','เทียบกับเกณฑ์โดยไม่แทนการตรวจอื่น'],['poiCode','string หรือ null','รหัสที่อยู่ในรายการอนุญาต'],['speechAct','ค่าจากชุดที่กำหนด','แยกคำขอ คำถาม ปฏิเสธ สมมติ และข้อความยกมา'],['needsClarification','boolean','แจ้งว่าต้องระบุจุดหมายเพิ่ม'],['decisionSource','rules หรือ model','Backend เพิ่มเพื่ออธิบายเส้นทางที่เลือก']])
p('ตัวอย่างผลที่ยอมรับได้คือ {"intent":"navigate_to","confidence":0.95,"poiCode":"<UUID ที่มีอยู่จริง>","speechAct":"request","needsClarification":false} ระบบไม่ยอมรับ destination ที่เป็นฟิลด์อิสระแทน poiCode ไม่ยอมรับ UUID ที่แต่งขึ้น และตรวจว่าการกลับแท่นไม่ได้ถูกใช้แทนคำขอไป POI ที่ระบุชื่อไว้')
h('4.6 เว็บไซต์ผู้ช่วยและการสนทนา')
p('เว็บไซต์ผู้ช่วยเปิดได้ที่ localhost:3001/ai ขณะที่หน้า AI ของ Admin ใช้ดูสถานะและบันทึก ช่องทาง Kiosk ใช้ Backend ตรวจการเลือกจุดหมายและสิทธิ์ก่อนเริ่มงาน การถามตำแหน่งคืนค่าจาก POI โดยไม่สร้างงาน เว็บไซต์แชตแสดงคำตอบทยอยมาและแหล่งข้อมูลเมื่อใช้การค้นหา')
p('ระบบมีการส่งข้อความสนทนาช่วงสั้นสำหรับแชตทั่วไป แต่การเชื่อมคำตอบ เช่น “ห้อง A” เข้ากับคำถามจุดหมายที่ค้างยังไม่ได้พัฒนา จึงไม่ควรสรุปว่าจำบริบทนำทางได้แล้ว กรณี “ห้องนั้น” ที่ไม่มีชื่ออ้างอิงชัดเจนในข้อความปัจจุบันจะถามให้ระบุใหม่')
h('4.7 แนวทางติดตั้งและการจัดการค่ากำหนด')
p('เริ่มโหมดพัฒนาด้วย Node.js และ pnpm ตาม packageManager ในโปรเจ็กต์ ติดตั้ง dependencies แล้วคัดลอกไฟล์ตัวอย่างค่ากำหนด ตั้ง DB_MODE=memory และ ROBOT_MODE=sim สำหรับการทดลองที่ไม่ต่อหุ่น หากใช้ PostgreSQL ให้ตั้งฐานข้อมูลและ migration ตามคู่มือ phase1-setup การเปิด AI ใช้ AI_MODE=local และปลายทาง llama-server ตามคู่มือ local AI')
p('รายงานไม่บรรจุรหัสผ่าน SSH, token, session secret หรือค่า .env จริง ที่อยู่เครือข่ายและบัญชีอุปกรณ์ควรกำหนดในสภาพแวดล้อมติดตั้ง การเปิดระบบใหม่ต้องตรวจโหมดจำลอง/จริงบนหน้าจอก่อนส่งคำสั่ง คู่มือลำดับการเริ่มและการทดสอบอยู่ในภาคผนวก ก')

page('บทที่ 5 ผลการดำเนินงานและการทดสอบ')
h('5.1 หลักฐานและช่วงเวลาของผลทดสอบ')
p('บทนี้ใช้ผล benchmark JSON ที่บันทึกวันที่ 7 ตุลาคม 2569 และผลทดสอบที่บันทึกใน PROGRESS.md ไม่ได้รันหุ่นหรือโมเดลใหม่ระหว่างจัดทำรูปเล่มวันที่ 12 ตุลาคม 2569 ผลเก่าที่อ่านฮาร์ดแวร์ได้จึงเป็นหลักฐานความสามารถในช่วงที่ทดสอบ ไม่ใช่คำยืนยันว่าหุ่นเชื่อมต่ออยู่ในปัจจุบัน')
p('เอกสารภาพรวมบางไฟล์มีสถานะเก่าก่อนเพิ่ม AI กับหุ่นจริง จึงใช้โค้ดปัจจุบัน รายงาน benchmark และบันทึกงานล่าสุดร่วมกัน แยกว่ามีโค้ดส่งงานได้ แต่ผู้ช่วยยังไม่ได้ตรวจการถึงจุดหมายหรือการหยุดทางกายภาพ หลักฐานทดสอบ API จำลองไม่ถือเป็นผลตรวจรับอุปกรณ์')
h('5.2 วิธีการทดสอบ benchmark')
p('ชุดข้อมูลมี 40 ข้อความและ 6 POI จำลอง รวมชื่อเรียกซ้ำและจุดส่วนตัว เปรียบเทียบ policy เดิม วิธี hybrid และวิธี model เป็นหลักที่ยังคงกฎหยุด/คำปฏิเสธ/สมมติ/ข้อความยกมา ทุกคำขอใช้ dryRun=true และบังคับ ROBOT_MODE=sim, DB_MODE=memory network allowlist จำกัดเฉพาะ origin ของโมเดล พร้อมตรวจว่าไม่มี job ถูกสร้าง')
p('กรณีผ่านต้องมี outcome ตรงกับ expected และรหัสเป้าหมายตรงเมื่อกรณีนั้นกำหนด target ตัวชี้วัด unsafeMotion ใน runner หมายถึงการเสนอ navigate/dock ทั้งที่ expected ไม่ใช่การเดิน หรือเลือกเป้าหมายผิด ไม่ครอบคลุมอันตรายทางกายภาพทุกชนิด จึงใช้คำว่า “เสนอเดินผิดเจตนา/จุดหมาย” ในรายงาน')
p('อัตราผ่านคำนวณจากจำนวนกรณีผ่านหารด้วยจำนวนกรณีทั้งหมดคูณ 100 ส่วน latency วัดตั้งแต่รับคำขอใน runner จนคืนผลที่ตรวจ policy แล้ว ใช้ nearest-rank สำหรับ p95 มีการปัดเป็นมิลลิวินาที ค่า 0 ms หมายถึงค่าต่ำกว่า 0.5 ms ไม่ใช่การประมวลผลทันทีโดยไม่มีเวลา')
h('5.3 ผลเปรียบเทียบการตีความคำสั่ง')
rows=[]
for name,label in [('baseline','ก่อนปรับปรุง'),('hybrid','กฎ + โมเดล'),('model','โมเดล + กฎป้องกัน')]:
    s=reports[name]['summary'];rows.append([label,f"{s['passed']}/{s['total']}",f"{s['accuracyPct']}%",s['unsafeMotion'],s['medianMs'],s['p95Ms']])
table('ตารางที่ 5.1 ผล regression benchmark บนพีซี',['วิธี','ผ่าน','อัตราผ่าน','เสนอเดินผิด','Median ms','p95 ms'],rows)
figure('benchmark-accuracy','ภาพที่ 5.1 อัตราผ่านของชุด regression เดียวกัน')
p('วิธี hybrid ผ่าน 40/40 โดยใช้กฎ 33 กรณีและโมเดล 7 กรณี p95 เท่ากับ 4,684 ms และค่าสูงสุด 6,184 ms median ที่ปัดเป็น 0 ms เกิดจากกรณีส่วนใหญ่ไม่ต้องรอ inference ผลนี้อธิบายการลดเวลารอสำหรับรูปประโยคที่กฎรองรับ ไม่ได้หมายความว่าโมเดลตอบทุกคำสั่งต่ำกว่า 1 ms')
p('วิธี model ผ่าน 40/40 โดยใช้กฎ 10 กรณีและโมเดล 30 กรณี รอบนี้มีการตรวจ Kiosk ร่วมช่วงต้น จึงอาจมีเวลารอคิวร่วมและไม่ควรสรุปเป็นการเปรียบเทียบความเร็วแบบควบคุมทุกตัวแปร นอกจากนี้ prompt cache อาจอุ่นและชุดข้อมูลเป็นชุด regression ที่ใช้พัฒนา ไม่ใช่ unseen test set')
h('5.4 ผลตามกลุ่มกรณีทดสอบ')
table('ตารางที่ 5.2 ผล hybrid ตามกลุ่มข้อความ',['กลุ่ม','จำนวน','ผ่าน','เสนอเดินผิด'],[[group,stats['total'],stats['passed'],stats['unsafeMotion']] for group,stats in reports['hybrid']['groups'].items()])
p('กรณีที่ต้องเข้าใจความหมายมีข้อความที่เอ่ยจุดหมายก่อนคำขอและประโยคอ้อม ส่วนกรณีป้องกันครอบคลุมคำปฏิเสธ คำสั่งสมมติ ข้อความยกมา คำถามความสามารถ หลายจุดหมาย จุดส่วนตัว และข้อความให้ละเลยกฎ ข้อมูลทุกกรณีอยู่ในภาคผนวก ข เพื่อให้ตรวจย้อนหลังและรันซ้ำได้')
h('5.5 ผลทดสอบ Backend และ Kiosk')
p('บันทึกผล Backend ผ่าน 51 tests ครอบคลุมสิทธิ์ API, WebSocket, แผนที่, POI, manual lock, การนำทาง, การหยุดผล AI ที่รออยู่ และ Kiosk ส่งงานผ่าน adapter จำลอง มีการตรวจ semantic output ที่ผิดรูปแบบและ UUID ที่แต่งขึ้น รวมทั้งหุ่นจำลองที่ไม่เก็บ aliases ต้องไม่รายงานว่าบันทึกสำเร็จ')
p('สคริปต์ตรวจ Kiosk กับ POI จริงในรอบก่อนหน้ารับคำขอ 4 สำนวนและเลือก UUID ของ “หน้าประตู” เดียวกันใน dry-run ได้แก่ เดินไปที่…, ช่วยพาไปที่…หน่อย, ไปส่งที่…ได้ไหมครับ และ …คือจุดหมายของผม คุณช่วยไปส่งผมที่นั่นได้ไหม การถามตำแหน่งคืนพิกัดจาก POI โดยไม่สร้างงาน ผลนี้ไม่ยืนยันการถึงจุดหมายจริง')
h('5.6 สิ่งที่ยังไม่มีผลตรวจรับ')
table('ตารางที่ 5.3 ขอบเขตที่ยังต้องตรวจเพิ่ม',['รายการ','หลักฐานปัจจุบัน','งานที่ต้องทำ'],[['เดิน/ถึง POI/กลับแท่น','โค้ดและการทดสอบ API จำลอง','ทดลองควบคุมจริงในพื้นที่ที่ผู้ใช้ดูแล'],['หยุดจริง/เครือข่ายหลุด','tests การยกเลิกและ manual watchdog','ตรวจพฤติกรรมอุปกรณ์'],['POI CRUD บนหุ่น','อ่าน POI จริง + tests จำลอง','ตรวจเขียน ลบ และการคงข้อมูลจริง'],['Jetson','เตรียม source/model และพัก build','build, deploy และ benchmark'],['STT/TTS','ยังไม่เปิดใช้','พัฒนาและวัดเสียง/ความหน่วง'],['สำรวจจริง','ตัวเลือกโหมดบนเว็บ','ต่อ API และตรวจรับการสร้างแผนที่'],['ภาษาใหม่และความพึงพอใจ','ยังไม่มีชุด unseen/แบบสำรวจ','เก็บข้อมูลและทดสอบแยกจากชุดพัฒนา']])

page('บทที่ 6 สรุปผลและแนวทางพัฒนาต่อ')
h('6.1 สรุปผลการพัฒนา')
p('โครงงานมีเว็บไซต์สำหรับเจ้าหน้าที่และผู้มาเยือน Backend กลาง ฐานข้อมูล และตัวเชื่อม Hermes สามารถแสดงข้อมูลแผนที่ ตำแหน่ง LiDAR และ POI ตามหลักฐานรอบก่อนหน้า มีคำสั่งควบคุมและนำทางที่ตรวจสิทธิ์และสถานะหุ่น การพัฒนา AI วิธีแรกและวิธีที่สองช่วยให้รองรับคำขอหลายสำนวนและชื่อเรียกอื่น พร้อมโมเดลตีความประโยคอ้อม')
p('ผล regression แสดงว่า policy ที่ปรับแล้วผ่านครบ 40 กรณีในชุดเดียวกัน และไม่เสนอเดินผิดเจตนาหรือผิดเป้าหมายในชุดนี้ แต่ยังไม่เพียงพอต่อการสรุปว่าระบบพร้อมใช้งานหน้างานทั้งหมด งานตรวจรับการเคลื่อนที่ ความพร้อม Jetson เสียง และบริบทนำทางยังคงเป็นข้อจำกัดที่รายงานไว้อย่างชัดเจน')
h('6.2 ปัญหาที่พบและแนวทางแก้ไข')
table('ตารางที่ 6.1 ปัญหาทางเทคนิคและแนวทางแก้',['ปัญหา','แนวทางในระบบ'],[['คำขอต้องใช้รูปประโยคเดิม','เพิ่มสำนวน ชื่อเรียกอื่น และ model fallback'],['รูปสระไทยหรือขอบเขตชื่ออังกฤษจับไม่ตรง','normalize NFC และตรวจขอบเขตคำ/ช่วงข้อความ'],['โมเดลคืน JSON ผิดรูปแบบ','กำหนด schema ในคำขอ/prompt และตรวจ Zod'],['ชื่อซ้ำหรือหลายจุดหมาย','ถามให้ระบุจุดเดียว ไม่เดารหัส'],['จุดหมายถูกแก้ขณะรอ AI','อ่าน POI ใหม่ก่อน dispatch'],['ผล AI กลับมาหลังสั่งหยุด','ตรวจ command epoch และ signal'],['firmware ไม่มีค่าความเร็ว','แสดงไม่ทราบแทนการคำนวณเดา'],['หุ่นออฟไลน์แต่ต้องเลือกโหมด','บันทึก preference บนเว็บโดยไม่ส่งคำสั่ง']])
h('6.3 แนวทางพัฒนาต่อ')
p('ลำดับถัดไปด้าน AI คือเก็บคำถามจุดหมายที่ค้างในขอบเขตบทสนทนา รองรับคำตอบ “ห้อง A” และการอ้าง “ที่นั่น” เฉพาะเมื่อมีบริบทชัดเจน ต้องหมดอายุ ล้างเมื่อยกเลิก และป้องกันสร้างงานซ้ำ โดยเพิ่ม benchmark หลายข้อความและการหยุดระหว่างถามกลับ')
p('ด้านโหมดสำรวจควรเริ่มจากผู้ใช้บังคับเดินและหุ่นสร้างแผนที่ พร้อมแยกโหมดที่เลือกกับโหมดที่อุปกรณ์ยืนยัน ก่อนพิจารณาสำรวจอัตโนมัติ ด้าน Jetson ต้องเริ่ม build ที่พัก ตรวจความเข้ากันได้ วัด RAM latency และระบบเสียงจริง ไม่เลือกโมเดลปลายทางจากผลพีซีเพียงอย่างเดียว')
p('การประเมินเพิ่มเติมควรมีชุดประโยคใหม่ที่แยกจากชุดพัฒนา ทดลองซ้ำหลายรอบภายใต้โหลดที่ควบคุม วัดความหน่วงและความสำเร็จในการเดินจริง ตรวจ Wi-Fi หลุด และเก็บความพึงพอใจผู้ใช้ด้วยวิธีที่อธิบายกลุ่มตัวอย่างได้ ก่อนสรุปผลประสิทธิภาพหรือประโยชน์เชิงบริการ')
h('6.4 รายการที่ต้องเติมก่อนส่งเล่ม')
for text in ['ชื่อโครงงานที่อนุมัติ ผู้จัดทำ รหัสนักศึกษา อาจารย์ที่ปรึกษา ภาควิชา คณะ มหาวิทยาลัย และปีการศึกษา','ข้อความกิตติกรรมประกาศและหน้าอนุมัติที่ตรวจโดยผู้จัดทำ','ภาพหน้าจอจริงและภาพอุปกรณ์ที่มีสิทธิ์ใช้ พร้อมคำบรรยายและวันทดสอบ','ผลตรวจรับการเคลื่อนที่จริง และผล Jetson/เสียงเฉพาะเมื่อมีหลักฐานใหม่','รูปแบบอ้างอิง รายละเอียดการเข้าเล่ม และรูปแบบสารบัญตามข้อกำหนดของหลักสูตร']:p(text)

page('บรรณานุกรม')
refs=[('[1] SLAMTEC.','Hermes REST API / Swagger documentation.','https://docs.slamtec.com/#/'),('[2] Qwen.','Qwen3.5-2B model card.','https://huggingface.co/Qwen/Qwen3.5-2B'),('[3] Socket.IO.','Delivery guarantees, version 4 documentation.','https://socket.io/docs/v4/delivery-guarantees/'),('[4] ggml-org.','llama.cpp server README, revision b10545.','https://github.com/ggml-org/llama.cpp/blob/b10545/tools/server/README.md')]
for author,title,url in refs:p(f'{author} {title} สืบค้น 12 ตุลาคม 2569. {url}')
for i,(name,desc) in enumerate([('SPEC.md','ข้อกำหนดและแผนงาน'),('PROGRESS.md','บันทึกความคืบหน้าและขอบเขตการตรวจ'),('apps/backend/prisma/schema.prisma','โครงสร้างฐานข้อมูล'),('apps/backend/src/services/robotIntent.ts','การตีความเจตนาแบบ structured output'),('apps/backend/src/services/robotLanguage.ts','กฎภาษาและการจับคู่ชื่อ'),('apps/backend/src/services/robotAiService.ts','การตรวจและส่งคำสั่ง AI'),('apps/backend/src/services/navigationPolicy.ts','ข้อกำหนดนำทาง'),('reports/ai-benchmark/hybrid.json','ผล benchmark hybrid วันที่ 7 ตุลาคม 2569')],5):p(f'[{i}] โครงการ Hermes Reception Control. {desc}. เอกสาร/โค้ดภายใน repository: {name}. ตรวจอ่าน 12 ตุลาคม 2569.')

page('ภาคผนวก ก คู่มือเริ่มต้นและ API')
h('ก.1 เริ่มระบบโดยไม่ต่อหุ่น')
for text in ['1. ติดตั้ง Node.js และ pnpm แล้วใช้ pnpm install ที่รากโปรเจ็กต์','2. คัดลอก .env.example เป็นไฟล์ค่ากำหนด ตั้ง DB_MODE=memory และ ROBOT_MODE=sim สำหรับการพัฒนา','3. ตั้งค่า Kiosk จากไฟล์ตัวอย่างโดยให้ token ฝั่ง proxy ตรงกับ Backend ไม่ใส่ token ในรายงาน','4. ใช้ pnpm dev เปิด Admin ที่ localhost:3000 และ Kiosk ที่ localhost:3001','5. เปิด localhost:3001/ai สำหรับแชต หรือหน้า Control ของ Admin สำหรับตัวเลือกโหมด','6. เมื่อต้องใช้โมเดล local อ่าน docs/phase2-local-ai.md และ scripts/start-local-ai.ps1 ก่อนเปิด AI_MODE=local']:p(text)
h('ก.2 คำสั่งตรวจสอบและ benchmark')
for command in ['pnpm typecheck','pnpm --filter @hermes/backend test','pnpm --filter @hermes/backend benchmark:ai -- --label hybrid','pnpm --filter @hermes/backend benchmark:ai -- --strategy model --label model','pnpm --filter @hermes/backend benchmark:ai -- --repeat 3 --label hybrid-repeat']:p(command)
p('คำสั่ง benchmark ต้องมี model server พร้อมใช้งาน แต่หุ่นต้องเป็นตัวจำลองตามที่ runner บังคับ สคริปต์คืน exit code 1 เมื่อมีกรณีไม่ผ่าน ผลรายงานชื่อเดิมถูกเขียนทับได้ จึงควรเลือก label ใหม่เมื่อเก็บผลคนละรอบ')
h('ก.3 API หลักของระบบ')
table('ตารางที่ ก.1 API หลักตาม routes',['Method / path','หน้าที่','การยืนยันตัวตน'],[['GET /api/robot/state','สถานะหุ่น','Admin session'],['GET /api/robot/live-map','แผนที่บนหุ่น','Admin session'],['POST /api/robot/lock','ถือสิทธิ์ควบคุม','operator/admin'],['POST /api/robot/manual','ควบคุมด้วยมือ','สิทธิ์และ manual lock'],['POST /api/robot/navigate','นำทางไปพิกัด','operator/admin'],['POST /api/robot/dock','กลับแท่น','operator/admin'],['POST /api/robot/estop','หยุดฉุกเฉิน','ผู้ใช้เข้าสู่ระบบ'],['POST/PUT/DELETE /api/robot/pois','จัดการ POI','operator/admin'],['POST /api/ai/command','คำสั่ง AI เจ้าหน้าที่','operator/admin'],['GET /api/kiosk/destinations','POI สาธารณะ','device token'],['POST /api/kiosk/intent','ตีความ/คำสั่ง รวม dry-run','device token'],['POST /api/kiosk/chat','แชตและคำสั่งผู้มาเยือน','device token'],['POST /api/kiosk/cancel','ยกเลิกงาน Kiosk','device token']])

page('ภาคผนวก ข ชุด benchmark และผลรายกรณี')
p('ข้อมูลต่อไปนี้อ่านจาก cases.json และผล hybrid.json ที่บันทึกไว้ รหัส target เป็นลำดับ POI ใน fixture ไม่ใช่รหัสฐานข้อมูลจริง กรณี reject/clarify ต้องไม่เริ่มเดิน')
table('ตารางที่ ข.1 POI จำลองใน benchmark',['ลำดับ','ชื่อ','ชื่อเรียกอื่น','การเข้าถึง'],[[i,poi['name'],', '.join(poi['aliases']),'ส่วนตัว' if poi.get('private') else 'สาธารณะ'] for i,poi in enumerate(dataset['catalog'],1)])
table('ตารางที่ ข.2 ข้อความ ผลที่คาดหวัง และผล hybrid',['กรณี / ข้อความ','Expected','Actual','ms'],[[row['id']+'\n'+row['text'],row['expected']+(f" / POI {row['target']}" if row.get('target') else ''),row['outcome']+' / '+('ผ่าน' if row['pass'] else 'ไม่ผ่าน'),row['elapsedMs']] for row in reports['hybrid']['rows']])
h('ข.1 การทำซ้ำและตรวจหลักฐาน')
p('Dataset SHA-256: '+reports['hybrid']['datasetSha256'])
p('ผล hybrid เริ่ม '+reports['hybrid']['startedAt']+' และจบ '+reports['hybrid']['finishedAt']+' (UTC) ข้อมูลเวลาต้นฉบับเก็บใน JSON ไฟล์รายงานผลไม่เก็บข้อมูลรับรองการเชื่อมต่ออุปกรณ์')

page('ภาคผนวก ค แนวทางตรวจรับและสถานะงาน')
table('ตารางที่ ค.1 รายการตรวจรับหน้างาน (ยังไม่ได้บันทึกผลใหม่)',['กรณี','วิธีตรวจเมื่อพร้อม','สถานะ'],[['นำทางถึง POI','ผู้ใช้ควบคุมพื้นที่และยืนยันจุดหมายก่อนเริ่ม','รอตรวจ'],['หยุด/ยกเลิก/E-Stop','ตรวจระหว่างงานตามขั้นตอนอุปกรณ์','รอตรวจ'],['กลับแท่น','ตรวจ home และแนวเข้าจอดกับสถานที่จริง','รอตรวจ'],['เครือข่ายหลุด','ตรวจ watchdog และสถานะหลังกลับมา','รอตรวจ'],['POI คงอยู่หลังเริ่มใหม่','อ่าน/แก้/ลบและตรวจหลังรีสตาร์ต','รอตรวจ'],['สำรวจและบันทึกแผนที่','เชื่อมโหมดจริงและยืนยันอ่านกลับ','ยังไม่พัฒนาเสร็จ'],['Jetson และเสียง','วัด latency, RAM, STT/TTS บนอุปกรณ์','ยังไม่ทดสอบ']])
p('แบบบันทึกแต่ละกรณีควรมีวันที่ ผู้ควบคุม firmware รุ่นแผนที่ จุดหมาย สภาพเครือข่าย ขั้นตอน ผลที่คาดหวัง ผลจริง และไฟล์หลักฐาน ไม่ควรเติมสถานะผ่านจากการทดสอบจำลองแทนผลฮาร์ดแวร์')
h('ค.1 ที่มาของรูปแบบเอกสาร')
p('ตรวจไฟล์ตัวอย่างทั้งหมดในโฟลเดอร์ตัวอย่างรูปเล่ม โดยยึดเว็บแอปพลิเคชันบอร์ดประชาสัมพันธ์ออนไลน์.docx และ TransportationManagementSystem.docx เป็นแนวทางเล่มไทย ตัวอย่าง PDF หลายเล่มใช้โครงสร้าง 6 บท ได้แก่ บทนำ แนวคิด/เทคโนโลยี วิธีดำเนินงาน การพัฒนา ผลดำเนินงาน และบทสรุป จึงใช้โครงสร้าง 6 บทในร่างนี้')
p('ใช้กระดาษ A4 ระยะขอบบน/ซ้าย 1.5 นิ้ว ล่าง/ขวา 1 นิ้ว ตามค่าที่พบในตัวอย่าง และใช้ TH Sarabun New 16 pt ที่มีในเครื่องแทนชื่อฟอนต์ TH SarabunPSK ของบางตัวอย่าง ไม่คัดลอกเนื้อหาหรือชื่อบุคคลจากตัวอย่างมาใช้กับโครงงานนี้')
h('ค.2 รายการภาพและตาราง')
for caption in figures:p(caption)
for caption in tables:p(caption)

settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.save(OUT/'AI-Robot-Butler-Report.docx')
(OUT/'AI-Robot-Butler-Report-source.md').write_text('\n\n'.join(text_source)+'\n',encoding='utf-8')
manifest={'date':'2026-10-12','status':'draft; personal/institutional details need completion','templateCount':len(template_info),'chapters':6,'figures':len(figures),'tables':len(tables),'sources':{str(path.relative_to(ROOT)):hashlib.sha256(path.read_bytes()).hexdigest() for path in [ROOT/'SPEC.md',ROOT/'PROGRESS.md',ROOT/'benchmarks/robot-ai/cases.json',*[ROOT/f'reports/ai-benchmark/{name}.json' for name in reports]]},'hardwareAccess':'none; report preparation only'}
(OUT/'report-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUT),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'templates':len(template_info)},ensure_ascii=False))

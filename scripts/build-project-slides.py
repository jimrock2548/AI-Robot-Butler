"""Editable native PowerPoint + matching vector PDF. No hardware/model access."""
import sys, json, html, math, shutil, hashlib, zipfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.tools/report-python'))
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
import pymupdf
from PIL import Image, ImageOps, ImageDraw

OUT=ROOT/'reports/presentation';OUT.mkdir(parents=True,exist_ok=True)
FONTS=ROOT/'.tools/presentation-fonts';FONTS.mkdir(parents=True,exist_ok=True)
font_source=Path('C:/Users/admin/AppData/Local/Microsoft/Windows/Fonts')
for source,target in [('THSarabunNew.ttf','regular.ttf'),('THSarabunNew Bold.ttf','bold.ttf')]:shutil.copyfile(font_source/source,FONTS/target)
archive=pymupdf.Archive(str(FONTS))
CSS="""@font-face{font-family:Sarabun;src:url(regular.ttf)}@font-face{font-family:Sarabun;src:url(bold.ttf);font-weight:700}*{margin:0;padding:0}body{font-family:Sarabun}p{margin:0}div{white-space:normal}a{color:inherit;text-decoration:none}"""
W,H=960,540
BG='F7F8F5';INK='142D3B';MUTED='526671';TEAL='087F80';LIGHT='E4F1EF';ORANGE='B95A26';WHITE='FFFFFF';LINE='D8E2E4';SOFT='EDF1F2'
prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
prs.core_properties.title='AI Robot Butler — Project Presentation';prs.core_properties.author='[กรอกชื่อผู้จัดทำ]'
pdf=pymupdf.open();slides=[];current=None;page=None
def color(value):return tuple(int(value[i:i+2],16)/255 for i in (0,2,4))
def rect(x,y,w,h,fill=WHITE,stroke=None,rounded=True):
    shape=current.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
    shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(fill)
    if stroke:shape.line.color.rgb=RGBColor.from_string(stroke);shape.line.width=Pt(.8)
    else:shape.line.fill.background()
    if rounded:shape.adjustments[0]=.1
    page.draw_rect(pymupdf.Rect(x,y,x+w,y+h),color=color(stroke) if stroke else None,fill=color(fill),radius=.1 if rounded else None,width=.8)
    return shape
def tx(text,x,y,w,h,size=26,bold=False,fill=INK,align='left'):
    shape=current.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));tf=shape.text_frame
    tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(text.split('\n')):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph();para.text=line;para.font.name='TH Sarabun New';para.font.size=Pt(size);para.font.bold=bold;para.font.color.rgb=RGBColor.from_string(fill)
        for tag in ['a:ea','a:cs']:
            node=OxmlElement(tag);node.set('typeface','TH Sarabun New');para.font._rPr.append(node)
        para.font._rPr.set('lang','th-TH')
        para.line_spacing=1.08;para.space_after=Pt(0);para.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
    markup=f'<div style="font-size:{size}pt;line-height:1.08;color:#{fill};font-weight:{700 if bold else 400};text-align:{align}">'+html.escape(text).replace('\n','<br/>')+'</div>'
    spare,scale=page.insert_htmlbox(pymupdf.Rect(x,y,x+w,y+h),markup,css=CSS,archive=archive,scale_low=1)
    if spare<0:raise ValueError(f'Text overflow on slide {len(slides)}: {text[:70]} ({w}x{h}, {size}pt)')
    slides[-1]['text'].append(text)
    return shape
def arrow(x1,y1,x2,y2,fill=MUTED):
    shape=current.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2));shape.line.color.rgb=RGBColor.from_string(fill);shape.line.width=Pt(1.8)
    from pptx.oxml.xmlchemy import OxmlElement
    end=OxmlElement('a:tailEnd');end.set('type','triangle');shape._element.spPr.get_or_add_ln().append(end)
    page.draw_line((x1,y1),(x2,y2),color=color(fill),width=1.8)
    angle=math.atan2(y2-y1,x2-x1);p1=(x2-9*math.cos(angle-.4),y2-9*math.sin(angle-.4));p2=(x2-9*math.cos(angle+.4),y2-9*math.sin(angle+.4))
    page.draw_polyline([(x2,y2),p1,p2,(x2,y2)],color=color(fill),fill=color(fill))
def new(title,kicker,notes,source='',dark=False):
    global current,page
    current=prs.slides.add_slide(prs.slide_layouts[6]);page=pdf.new_page(width=W,height=H)
    slides.append({'number':len(slides)+1,'title':title,'notes':notes,'source':source,'text':[]})
    current.background.fill.solid();current.background.fill.fore_color.rgb=RGBColor.from_string(INK if dark else BG)
    page.draw_rect(page.rect,color=None,fill=color(INK if dark else BG))
    fg=WHITE if dark else INK;dim='AFC6CD' if dark else MUTED
    rect(48,35,34,5,TEAL,rounded=False);tx(kicker,92,23,790,27,17,True,dim)
    tx(title,48,63,864,69,36,True,fg)
    page.draw_line((48,489),(912,489),color=color('35515E' if dark else LINE),width=.7)
    tx('AI ROBOT BUTLER  ·  PROJECT DRAFT',48,503,540,22,14,False,dim)
    tx(f'{len(slides):02d} / 18',807,501,104,24,16,False,dim,'right')
    current.notes_slide.notes_text_frame.text=notes+('\n\nหลักฐาน: '+source if source else '')
def card(x,y,w,h,number,title,body,accent=TEAL):
    rect(x,y,w,h,WHITE,LINE);tx(number,x+18,y+16,w-36,36,25,True,accent)
    tx(title,x+18,y+60,w-36,58,29,True)
    tx(body,x+18,y+128,w-36,h-143,24)
def box(x,y,w,h,title,body,fill=WHITE,accent=INK):
    rect(x,y,w,h,fill,LINE);tx(title,x+12,y+10,w-24,34,25,True,accent);tx(body,x+12,y+52,w-24,h-60,21)
def strip(text,y=433,fill=LIGHT,ink=TEAL):rect(48,y,864,39,fill,rounded=True);tx(text,62,y+5,836,30,21,True,ink)

new('ระบบหุ่นยนต์ผู้ช่วยและนำทางด้วย AI','PROJECT PRESENTATION · 12 OCT 2026',
    'แนะนำชื่อโครงงานและผู้จัดทำ ระบบนี้เชื่อมเว็บผู้ดูแลและเว็บผู้มาเยือนเข้ากับหุ่นผ่าน Backend และเพิ่ม AI local สำหรับคำขอภาษาธรรมชาติ เนื้อหาครั้งนี้อ้างอิงโค้ดและผลทดสอบที่บันทึกไว้ ไม่ได้ต่อหุ่นระหว่างจัดทำสไลด์ ตั้งใจใช้สไลด์หลัก 16 หน้าในเวลาประมาณ 10–15 นาที และอีก 2 หน้าเป็นข้อมูลสำรอง',dark=True)
tx('AI ROBOT\nBUTLER',48,154,510,158,67,True,WHITE)
tx('Web-based assistant · Hermes Pro Max\nLocal Qwen · ควบคุมผ่าน Backend',52,337,500,66,27,False,'AFC6CD')
tx('[กรอกชื่อผู้จัดทำ]  |  [กรอกอาจารย์ที่ปรึกษา]',52,441,570,32,22,False,'AFC6CD')
box(620,160,278,85,'WEB','Admin + Visitor Kiosk',fill=LIGHT)
box(620,265,278,85,'BACKEND','สิทธิ์ · คำสั่ง · ข้อมูล',fill=WHITE)
box(620,370,278,85,'ROBOT + LOCAL AI','ภาพรวมแนวคิดของระบบ',fill=LIGHT)
arrow(759,245,759,261,'AFC6CD');arrow(759,350,759,366,'AFC6CD')

new('ปัญหาที่โครงงานต้องการแก้','01 · PROBLEM',
    'ปัญหาแรกคือผู้มาเยือนต้องการข้อมูลและจุดหมายที่ถูกต้อง ปัญหาที่สองคือเจ้าหน้าที่ต้องเห็นแผนที่และสถานะงานร่วมกัน ส่วนการใช้ AI มีความเสี่ยงตีความคำถามหรือคำปฏิเสธเป็นคำสั่ง จึงออกแบบให้ AI เสนอเจตนา และ Backend ตรวจอีกชั้นก่อนส่งงาน นี่เป็นโจทย์การออกแบบ ไม่ใช่ผลสำรวจผู้ใช้ที่เก็บแล้ว','SPEC.md; robotAiService.ts')
card(48,155,272,265,'01','ข้อมูลและจุดหมาย','ผู้มาเยือนต้องเลือก\nจุดที่มีอยู่จริง\nและใช้คำขอที่คุ้นเคย')
card(344,155,272,265,'02','การควบคุมร่วมกัน','เจ้าหน้าที่ต้องเห็น\nแผนที่ สถานะ และงาน\nพร้อมสิทธิ์ที่ชัดเจน')
card(640,155,272,265,'03','ความหมายของภาษา','คำถาม คำปฏิเสธ\nและคำขอที่กำกวม\nต้องไม่ทำให้เดินผิด')
strip('แนวทาง: แยกการตีความภาษาออกจากการอนุมัติคำสั่งหุ่น')

new('เป้าหมาย: เว็บเดียวกันในระบบ สองกลุ่มผู้ใช้','02 · OBJECTIVES',
    'เป้าหมายประกอบด้วยเว็บไซต์แยกสำหรับเจ้าหน้าที่และผู้มาเยือน การแสดงแผนที่และจัดการ POI การรับคำขอหลายสำนวน และการทดสอบก่อนทดลองเคลื่อนที่จริง จุดเน้นคือใช้รายการจุดหมายและข้อมูลจากหุ่นเป็นแหล่งอ้างอิง รวมทั้งตรวจงานและสิทธิ์ก่อนเริ่ม','PROGRESS.md; SPEC.md')
items=[('เว็บควบคุม + เว็บผู้ช่วย','Admin และ Kiosk ใช้ Backend ร่วมกัน'),('แผนที่และจุดหมาย','อ่าน SLAM / LiDAR และจัดการ POI'),('คำขอภาษาธรรมชาติ','กฎภาษา + aliases + โมเดล local'),('ตรวจผลก่อนเดิน','หุ่นจำลอง · dry-run · benchmark')]
for i,(a,b) in enumerate(items):
    y=151+i*77;rect(48,y,864,66,WHITE,LINE);tx(f'0{i+1}',65,y+10,56,42,29,True,TEAL);tx(a,136,y+6,335,49,28,True);tx(b,485,y+10,405,43,24)

new('สถาปัตยกรรม: Backend เป็นตัวกลาง','03 · SYSTEM ARCHITECTURE',
    'ลูกค้าสองเว็บไซต์ไม่สั่งหุ่นจากเบราว์เซอร์โดยตรง Admin ใช้ session และ Kiosk ใช้ token ผ่าน server-side proxy Backend เชื่อม PostgreSQL หุ่น และโมเดล Qwen ผ่าน llama-server โดยการค้นหา SearXNG เป็นบริการเสริม เส้นทางนี้ช่วยรวมการตรวจสิทธิ์ งาน และการบันทึกไว้ส่วนกลาง','apps/admin; apps/kiosk; apps/backend; packages/shared')
box(48,169,224,112,'ADMIN :3000','Session · ควบคุม / ดูงาน')
box(48,333,224,112,'KIOSK :3001','Proxy · device token')
box(356,225,250,171,'BACKEND :4000','Express + Socket.IO\nตรวจสิทธิ์และงาน\nPrisma / adapter',LIGHT,TEAL)
box(705,153,207,88,'POSTGRESQL','ข้อมูลและบันทึก')
box(705,260,207,88,'HERMES','แผนที่ / คำสั่ง')
box(705,367,207,88,'LOCAL QWEN','llama-server :8080')
arrow(272,225,354,270);arrow(272,389,354,342);arrow(606,260,703,197);arrow(606,309,703,304);arrow(606,361,703,411)

new('แยกประสบการณ์ของเจ้าหน้าที่กับผู้มาเยือน','04 · USERS & PERMISSIONS',
    'เจ้าหน้าที่มีบทบาท viewer, operator และ admin ส่วนผู้มาเยือนใช้ Kiosk และเลือกได้เฉพาะ POI ที่อนุญาต จุดที่ metadata.is_public เป็น false ถูกกรองออก Kiosk หยุดงานของตนเองได้ แต่หยุดงานที่มาจากเจ้าหน้าที่ไม่ได้ ข้อจำกัดตรวจใน Backend ไม่ได้อาศัยปุ่มที่ซ่อนบนหน้าจออย่างเดียว','routes/kiosk.ts; middleware/auth.ts; kioskHardwareAi.test.ts')
card(48,153,418,269,'ADMIN / OPERATOR','เว็บไซต์เจ้าหน้าที่','ดูแผนที่และสถานะ\nควบคุมหุ่นและจัดการ POI\nดูงาน ผู้ใช้ และบันทึก AI')
card(494,153,418,269,'VISITOR KIOSK','เว็บไซต์ผู้ช่วย','เลือกจุดหมายสาธารณะ\nพิมพ์ถามและขอให้นำทาง\nยกเลิกงานของ Kiosk')
strip('viewer ดูข้อมูลได้ · Kiosk หยุดงานของเจ้าหน้าที่ไม่ได้')

new('แผนที่สดและ POI จากหุ่น','05 · MAP & DESTINATIONS',
    'แผนที่กริดมี origin และ resolution ซึ่งต้องรักษาไว้เมื่อแสดงบนเว็บ รวมถึงค่าจุดกำเนิดติดลบ ระบบแสดงหุ่น แท่นชาร์จ LiDAR และ POI และรองรับซูมลาก การแก้ชื่อ พิกัด ทิศและ aliases ต้องอ่านข้อมูลกลับก่อนรายงานสำเร็จ ภาพบนสไลด์เป็นภาพอธิบาย ไม่ใช่ภาพแผนที่จริง ผลอ่านข้อมูลสดเป็นหลักฐานจากรอบที่เคยเชื่อมต่อ','LiveRobotMap.tsx; RobotMapEditor.tsx; hermesClient.ts; PROGRESS.md')
rect(48,158,398,278,WHITE,LINE)
for r in range(7):
    for c in range(10):rect(73+c*34,180+r*30,30,26,LINE if (r==1 and c<6) or (c==8 and r>3) else LIGHT,rounded=False)
rect(141,300,28,28,TEAL);tx('R',142,297,26,32,25,True,WHITE,'center')
for x,y,label in [(311,240,'POI'),(210,360,'DOCK')]:rect(x,y,63,27,ORANGE);tx(label,x+2,y,59,28,20,True,WHITE,'center')
tx('ภาพอธิบาย · ไม่ใช่แผนที่จริง',67,441,370,30,18,False,MUTED)
tx('ข้อมูลบนเว็บ',489,159,407,42,29,True)
tx('แผนที่ / POI / LiDAR ประมาณ 1 วินาที\nตำแหน่งประมาณ 0.5 วินาที\nHTTP fallback ประมาณ 2 วินาที',489,216,407,118,25)
tx('เพิ่ม · แก้ไข · ลบ POI\nตรวจชื่อ พิกัด ทิศ และ aliases ด้วย readback',489,351,407,91,25)

new('ก่อนนำทาง ต้องผ่านข้อกำหนดของ Backend','06 · ROBOT CONTROL',
    'การสั่งด้วย joystick ต้องถือ manual lock ที่มี TTL และ heartbeat การนำทางต้องมีแผนที่อายุไม่เกิน 5 วินาที อยู่ในโหมดระบุตำแหน่ง และมีคุณภาพอย่างน้อย 50 จุดหมายต้องอยู่ในแผนที่และเป็นพื้นที่ว่างตามค่าที่อ่าน หุ่นต้องไม่มีงานหรือ lock ขัดกัน มี halt และ E-Stop แต่พฤติกรรมหยุดจริงยังรอตรวจรับ โหมดกลับแท่นอัตโนมัติปิดไว้ในห้องแลป','navigationPolicy.ts; driver.ts; manualHardware.test.ts')
card(48,154,272,267,'MAP','ข้อมูลพร้อม','แผนที่อายุ ≤ 5 วินาที\nปิดโหมดสร้างแผนที่\nคุณภาพตำแหน่ง ≥ 50')
card(344,154,272,267,'TARGET','จุดหมายถูกต้อง','ใช้รหัส POI ที่มีจริง\nพิกัดอยู่ในแผนที่\nอยู่ในเซลล์พื้นที่ว่าง')
card(640,154,272,267,'CONTROL','สิทธิ์ไม่ขัดกัน','ตรวจงานและ manual lock\nมีคำสั่งหยุด / E-Stop\nไม่เริ่มเดินตอน startup')
strip('เงื่อนไขซอฟต์แวร์ผ่าน ≠ ตรวจรับการเดินจริงผ่าน',fill='F9ECE3',ink=ORANGE)

new('AI ใช้สองวิธีร่วมกัน','07 · NATURAL LANGUAGE',
    'วิธีแรกคือกฎรูปประโยคและชื่อเรียกอื่น เช่น ทางเข้าตรงกับ POI หน้าประตู ช่วยให้คำขอชัดเจนตอบเร็ว วิธีที่สองใช้โมเดลตีความประโยคอ้อมหรือจุดหมายที่อยู่ก่อนคำขอ โมเดลคืน JSON โดยจำกัดรหัสปลายทาง ความมั่นใจที่โมเดลรายงานไม่ใช่สิทธิ์ให้เดิน Backend ยังตรวจความหมายและสถานะหุ่น','robotLanguage.ts; robotIntent.ts; docs/benchmark-ai.md')
card(48,151,418,277,'METHOD 01','กฎ + ชื่อเรียกอื่น','“ช่วยพาไปหน้าประตูหน่อย”\n“ไปส่งที่ทางเข้าได้ไหมครับ”\n\nจับคำขอและ POI ที่ชัดเจน')
card(494,151,418,277,'METHOD 02','โมเดลตีความความหมาย','“หน้าประตูคือจุดหมายของผม\nคุณช่วยไปส่งผมที่นั่นได้ไหม”\n\nคืน intent / POI / speechAct')
strip('ระบบใช้งาน: กฎก่อน → โมเดลเมื่อจำเป็น → Backend ตรวจอีกครั้ง')

new('AI เสนอคำสั่ง แต่ไม่ได้เลือกพิกัดอิสระ','08 · COMMAND PIPELINE',
    'อธิบายเส้นทางจากข้อความถึงงาน: อ่านรายการ POI ที่อนุญาต จำแนกเจตนา ตรวจ JSON schema ตรวจชื่อและความกำกวม แล้วอ่านพิกัดใหม่ก่อน dispatch การปฏิเสธหรือถามกลับไม่สร้างงาน dry-run คืนเพียงผลตีความ ส่วนงานจริงผ่านเงื่อนไขการนำทาง คำหยุดที่รองรับข้ามโมเดล และทำให้ผลที่กำลังรอ dispatch ภายหลังไม่ได้','robotAiService.ts; robotIntent.ts; robotAi.test.ts')
for x,title,body in [(48,'ข้อความผู้ใช้','ไทย / อังกฤษ'),(272,'จำแนกเจตนา','กฎ หรือโมเดล'),(496,'ตรวจคำขอ','Schema + POI'),(720,'สร้างงาน','อ่านพิกัดใหม่')]:box(x,196,192,118,title,body,LIGHT if x==496 else WHITE)
for x in [240,464,688]:arrow(x,253,x+28,253)
tx('ปฏิเสธ / ถามกลับ',322,351,319,48,31,True,ORANGE,'center')
tx('ชื่อซ้ำ · ไม่พบชื่อ · คำปฏิเสธ · ข้อความยกมา',158,400,645,36,24,False,MUTED,'center')
arrow(593,316,497,348,ORANGE)
strip('dry-run แสดงผลอย่างเดียว · คำหยุดที่รองรับไม่ต้องรอโมเดล',y=446)

new('ประโยคคล้ายกัน อาจต้องทำคนละอย่าง','09 · INTENT EXAMPLES',
    'ใช้ตัวอย่างแสดงความต่างของคำขอ คำถาม คำปฏิเสธและความกำกวม ชื่อห้องประชุมหรือห้องเล็กในสไลด์เป็นตัวอย่างจาก benchmark ไม่ได้ยืนยันว่าบนหุ่นปัจจุบันมีจุดเหล่านี้ คำถามตำแหน่งต้องตอบข้อมูล คำปฏิเสธไม่เดิน และชื่อกำกวมต้องถามใหม่ การเก็บคำตอบถามกลับข้ามข้อความยังเป็นงานถัดไป','benchmarks/robot-ai/cases.json; robotLanguage.test.ts')
examples=[('ช่วยพาไปที่ห้องประชุมหน่อย','นำทางไป POI','request',TEAL),('ห้องประชุมอยู่ที่ไหน','ตอบตำแหน่ง','question',TEAL),('อย่าเพิ่งไปห้องประชุม','ไม่ส่งคำสั่งเดิน','negated',ORANGE),('ช่วยพาไปห้องเล็กหน่อย','ถามให้ระบุจุดเดียว','ambiguous',ORANGE)]
for i,(a,b,label,ink) in enumerate(examples):
    y=151+i*76;rect(48,y,864,65,WHITE,LINE);tx(a,65,y+12,477,40,27);tx(b,540,y+12,271,40,27,True,ink);tx(label,810,y+17,85,29,17,False,MUTED,'right')

new('โหมดปกติ / โหมดสำรวจ: ตั้งค่าได้ขณะออฟไลน์','10 · OPERATING MODES',
    'ตัวเลือกโหมดที่เพิ่มล่าสุดเป็น preference บนเว็บ บันทึกไว้ใน localStorage และเลือกได้เมื่ออุปกรณ์ไม่เชื่อมต่อ โหมดปกติหมายถึงแนวทางใช้แผนที่และ POI เดิม โหมดสำรวจเป็นตัวเลือกเตรียมสำรวจและสร้างแผนที่ แต่ยังไม่มี API เปลี่ยนโหมดจริง การเลือกหรือเชื่อมต่อใหม่ไม่ส่งคำสั่งและไม่เริ่มสำรวจอัตโนมัติ','RobotOperatingMode.tsx; PROGRESS.md 11 October 2026')
card(48,153,418,272,'NORMAL','โหมดปกติ','ใช้แผนที่และ POI เดิม\nสำหรับการนำทาง\n\nเลือกและจำตัวเลือกบนเว็บได้')
card(494,153,418,272,'EXPLORATION','โหมดสำรวจ','เตรียมสำรวจและสร้างแผนที่\nยังไม่เปิด mapping บนหุ่น\n\nการเชื่อมโหมดจริงเป็นงานถัดไป',accent=ORANGE)
strip('สถานะปัจจุบัน: เลือกโหมดบนเว็บ ≠ เปลี่ยนโหมดบนหุ่น',fill='F9ECE3',ink=ORANGE)

new('ออกแบบ benchmark ให้ทดสอบโดยไม่เดิน','11 · EVALUATION METHOD',
    'ชุด regression มี 40 ข้อความและ 6 POI fixture ทดสอบการจำแนกและ policy รวมกัน ให้ผ่านเมื่อ outcome และรหัสจุดหมายตรงกับ expected เปรียบเทียบ policy เดิม hybrid และ model เป็นหลัก Runner บังคับหุ่นจำลอง ฐานข้อมูล memory และ dry-run พร้อมจำกัด network เฉพาะโมเดล ไม่ใช้เป็นผลทดสอบการเดินจริง','benchmark-robot-ai.ts; cases.json; reports/ai-benchmark/*.json')
for x,big,label in [(48,'40','ข้อความทดสอบ'),(344,'6','POI จำลอง'),(640,'3','วิธีที่เปรียบเทียบ')]:
    rect(x,156,272,146,WHITE,LINE);tx(big,x+18,168,236,79,67,True,TEAL);tx(label,x+18,259,236,34,26)
tx('ครอบคลุม',48,326,170,39,29,True)
tx('คำขอสุภาพ · aliases · ประโยคอ้อม · กลับแท่น · หยุด\nคำถาม · คำปฏิเสธ · ชื่อกำกวม · จุดส่วนตัว',225,329,687,79,25)
strip('หุ่นจำลอง + memory + dry-run · ตรวจว่าไม่มี job ถูกสร้าง')

results={name:json.loads((ROOT/f'reports/ai-benchmark/{name}.json').read_text(encoding='utf-8')) for name in ['baseline','hybrid','model']}
new('ผล regression: จาก 21/40 เป็น 40/40','12 · BENCHMARK RESULTS',
    'ผลที่รายงานบันทึกวันที่ 7 ตุลาคม 2569 ไม่ใช่การรันใหม่ในวันจัดทำสไลด์ policy เดิมผ่าน 52.5% และมี 4 กรณีเสนอเดินผิดเจตนาหรือจุดหมาย วิธี hybrid และ model เป็นหลักผ่านครบ 40 ในชุดนี้ แต่ชุดเดียวกันนี้ใช้พัฒนาด้วย จึงยังไม่ใช่ unseen evaluation และไม่รับประกันทุกประโยค ผลทั้งหมดเป็นข้อเสนอใน dry-run ไม่มีการส่งหุ่นเดิน','baseline.json; hybrid.json; model.json — 7 October 2026')
for i,(key,label) in enumerate([('baseline','ก่อนปรับ'),('hybrid','กฎ + โมเดล'),('model','โมเดล + กฎป้องกัน')]):
    s=results[key]['summary'];y=161+i*80
    tx(label,48,y,215,37,27,True);rect(277,y+4,444,27,SOFT,rounded=False);rect(277,y+4,444*s['accuracyPct']/100,27,TEAL if key!='baseline' else '94A6AE',rounded=False)
    tx(f"{s['accuracyPct']:g}%",742,y-3,165,44,35,True,TEAL if key!='baseline' else MUTED)
rect(48,407,864,65,LIGHT);tx('เสนอเดินผิดเจตนา/จุดหมายในชุดนี้',63,415,474,39,26,True);tx('4 → 0 กรณี',649,408,245,56,38,True,TEAL,'right')
tx('PC · dry-run · ชุดที่ใช้พัฒนา ไม่ใช่ unseen test หรือผลการเดินจริง',48,369,864,28,18,False,MUTED)

new('หลักฐานที่ตรวจแล้ว และขอบเขตของผล','13 · VERIFICATION',
    'บันทึก Backend ผ่าน 51 tests และ typecheck ของ Admin Kiosk Backend ผ่าน มีการตรวจ Kiosk กับ POI จริงแบบ dry-run 4 สำนวนและเลือก UUID เดียวกัน การส่งงานถึง MoveTo ทดสอบผ่าน adapter จำลอง ส่วนการอ่านแผนที่และ LiDAR เป็นข้อมูลจากฮาร์ดแวร์ในรอบก่อนหน้า ไม่มีผลตรวจถึงจุดหมาย หยุดจริง หรือกลับแท่นจริงจากผู้ช่วย','PROGRESS.md; robotAi.test.ts; kioskHardwareAi.test.ts; check-kiosk-ai.mjs')
card(48,154,272,267,'51 TESTS','Backend','สิทธิ์และงาน\nPOI / schema / ยกเลิก AI\nทดสอบกับ adapter จำลอง')
card(344,154,272,267,'TYPECHECK','สามส่วนหลัก','Admin · Kiosk · Backend\nผ่านตามบันทึกโครงการ\nไม่ใช่ผลตรวจฮาร์ดแวร์')
card(640,154,272,267,'KIOSK DRY-RUN','POI จริง','คำขอ 4 สำนวน\nเลือก UUID เดียวกัน\nโดยไม่สร้างงานเดิน')
strip('ผลอ่านฮาร์ดแวร์เป็นหลักฐานรอบก่อน · รอบจัดทำสไลด์หุ่นไม่ได้เชื่อมต่อ')

new('ข้อจำกัดปัจจุบัน และงานที่ควรทำต่อ','14 · LIMITATIONS & ROADMAP',
    'ปัจจุบัน AI ทำงานบนพีซี Jetson เตรียมอุปกรณ์และโมเดลแล้วแต่ build พักและไม่มี benchmark บนอุปกรณ์ ระบบเสียงยังไม่เปิด บริบทถามกลับเช่นห้อง A หรือที่นั่นข้ามข้อความยังไม่ทำ โหมดสำรวจเป็น preference ฝั่งเว็บ ขั้นต่อไปควรเพิ่มการถามกลับแบบมีอายุ ทดสอบหลายข้อความ แล้วเชื่อมโหมดจริงและตรวจรับการเดินในพื้นที่ที่ผู้ใช้ควบคุม','PROGRESS.md; SPEC.md; docs/jetson-ai.md')
box(48,151,418,271,'ยังไม่มีผลยืนยัน','• การเดินถึง POI / หยุด / กลับแท่นจริง\n• benchmark และระบบเสียงบน Jetson\n• ความแม่นยำกับประโยคใหม่\n• โหมดสำรวจจริงผ่าน API',WHITE,ORANGE)
box(494,151,418,271,'ลำดับงานถัดไป','1. บริบทถามกลับ + benchmark หลายข้อความ\n2. เชื่อมโหมดจริงและอ่านกลับยืนยัน\n3. ตรวจรับการเดินกับผู้ควบคุม\n4. วัด Jetson / STT / TTS',LIGHT,TEAL)
strip('AI local บน PC ทำแล้ว · Jetson และเสียงยังไม่พร้อมใช้งานหน้างาน',fill='F9ECE3',ink=ORANGE)

new('เว็บ + ข้อมูลจริง + AI ที่ผ่านการตรวจคำขอ','15 · SUMMARY / Q&A',
    'สรุปว่าโครงงานรวมเว็บผู้ดูแล จอผู้ช่วย ข้อมูลแผนที่ และ AI local โดยตรวจคำขอผ่าน Backend ก่อนสร้างงาน ผล regression ดีขึ้นในชุดที่ทดสอบ แต่ยังไม่ปิดงานตรวจรับหน้างาน เปิดให้กรรมการถามเรื่องสถาปัตยกรรม policy ความกำกวม ความเร็ว หรือ Jetson ใช้สไลด์สำรอง 17 และ 18 ประกอบคำตอบได้',dark=True)
tx('แยกผู้ใช้  ·  ใช้ POI ที่มีจริง\nตรวจสิทธิ์และสถานะก่อนส่งงาน',52,169,830,104,43,True,WHITE)
tx('กฎ + โมเดล: ผ่าน 40/40 ในชุด regression\nขั้นถัดไป: บริบทสนทนาและตรวจรับฮาร์ดแวร์',52,302,830,84,30,False,'AFC6CD')
tx('ขอบคุณครับ / ค่ะ',52,422,530,44,35,True,WHITE);tx('Q & A',744,411,160,66,45,True,'72D3C5','right')

new('ข้อมูลสำรอง: ความหน่วงและข้อจำกัดการวัด','BACKUP A · BENCHMARK DETAILS',
    'ตอบคำถามเรื่องเวลา: hybrid ใช้กฎ 33 กรณีและโมเดล 7 กรณี จึง median ที่ปัดเป็น 0 ms หมายถึงต่ำกว่า 0.5 ms ส่วน p95 ประมาณ 4.684 วินาที วิธี model ใช้โมเดล 30 กรณีและกฎ 10 กรณี รอบนั้นมี Kiosk ร่วมช่วงต้น จึงมี queue ที่อาจกระทบเวลา ไม่ควรเปรียบเทียบประสิทธิภาพแบบควบคุมทุกตัวแปร ไม่ได้วัด RAM peak, TTFT, Jetson หรือการเดินจริง','reports/ai-benchmark/{baseline,hybrid,model}.json; docs/benchmark-ai.md')
head=['วิธี','Median','p95','ใช้โมเดล'];xs=[61,365,547,738];widths=[290,164,171,157]
rect(48,156,864,51,INK)
for x,w,label in zip(xs,widths,head):tx(label,x,164,w,37,25,True,WHITE)
for i,(key,label) in enumerate([('baseline','ก่อนปรับ'),('hybrid','กฎ + โมเดล'),('model','โมเดล + กฎป้องกัน')]):
    y=208+i*57;s=results[key]['summary'];rect(48,y,864,54,WHITE if i!=1 else LIGHT,LINE,False)
    values=[label,'< 1 ms' if key=='hybrid' else f"{s['medianMs']/1000:.3f} s",f"{s['p95Ms']/1000:.3f} s",'ไม่บันทึก source' if key=='baseline' else f"{s['model']}/40"]
    for x,w,value in zip(xs,widths,values):tx(value,x,y+10,w,36,24,False,INK)
tx('PC: i7-10700F · RAM รวม 63.9 GiB · Qwen3.5-2B Q5_K_M',48,397,864,31,21)
tx('Regression ที่ใช้พัฒนา · cache อาจอุ่น · รอบ model มี Kiosk ร่วมช่วงต้น',48,436,864,31,20,False,ORANGE)

new('ข้อมูลสำรอง: เอกสารและไฟล์อ้างอิง','BACKUP B · SOURCES',
    'แหล่งข้อมูลหลักคือโค้ดและหลักฐานของโครงการ รายงานรูปเล่มและ JSON benchmark สามารถเปิดตรวจรายกรณีได้ เอกสารผู้ผลิตและ model card ใช้เป็นแหล่งเทคโนโลยีประกอบ ไม่ใช้ชื่อบุคคลหรือผลการทดลองจากเล่มตัวอย่าง หากกรรมการต้องการผลการเดินจริงให้แจ้งว่ายังเป็นรายการรอตรวจรับ','reports/project-documentation; PROGRESS.md; repository source')
tx('หลักฐานในโปรเจ็กต์',48,154,417,43,30,True,TEAL)
tx('SPEC.md / PROGRESS.md\nreports/ai-benchmark/*.json\nbenchmarks/robot-ai/cases.json\napps/backend/src/services/\nreports/project-documentation/',48,214,427,173,24)
tx('เอกสารเทคโนโลยี',503,154,409,43,30,True,TEAL)
tx('SLAMTEC REST / Swagger\ndocs.slamtec.com\nQwen3.5-2B model card\nhuggingface.co/Qwen/Qwen3.5-2B\nllama.cpp server README · b10545',503,214,409,173,23)
strip('เนื้อหา ณ 12 ต.ค. 2569 · benchmark อ้างผลวันที่ 7 ต.ค. 2569')

assert len(slides)==18
deck=OUT/'AI-Robot-Butler-Presentation.pptx';pdf_path=OUT/'AI-Robot-Butler-Presentation.pdf'
prs.save(deck);pdf.save(pdf_path,garbage=4,deflate=True)
notes=['# โน้ตช่วยพรีเซนต์ AI Robot Butler','สไลด์หลัก 1–16 ประมาณ 10–15 นาที; สไลด์ 17–18 สำรองสำหรับตอบคำถาม','เติมข้อมูลผู้จัดทำและอาจารย์ในสไลด์ 1 ก่อนนำเสนอ']
for slide in slides:notes.extend([f"## {slide['number']}. {slide['title']}",slide['notes'],'หลักฐาน: '+slide['source'] if slide['source'] else ''])
(OUT/'speaker-notes.md').write_text('\n\n'.join(notes)+'\n',encoding='utf-8')
(OUT/'slides-content.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding='utf-8')
with zipfile.ZipFile(deck) as z:assert z.testzip() is None
check=Presentation(deck);assert len(check.slides)==18
assert all(slide.has_notes_slide for slide in check.slides)
for s in check.slides:
    for shape in s.shapes:
        assert shape.left>=0 and shape.top>=0
        assert shape.left+shape.width<=prs.slide_width and shape.top+shape.height<=prs.slide_height
with pymupdf.open(pdf_path) as d:
    assert len(d)==18
    for i,slide in enumerate(slides):
        # All PDF objects are bounded by the page; text renderer forbids overflow.
        assert len(d[i].get_text().strip())>80
        if i in [0,3,7,12,14,16]:d[i].get_pixmap(matrix=pymupdf.Matrix(1.25,1.25)).save(OUT/f'preview-{i+1:02d}.png')
    thumbs=[]
    for p in d:
        pix=p.get_pixmap(matrix=pymupdf.Matrix(.4,.4));im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);thumbs.append(im)
    canvas=Image.new('RGB',(384*3,246*6),'#dde4e7');draw=ImageDraw.Draw(canvas)
    for i,im in enumerate(thumbs):x=(i%3)*384;y=(i//3)*246;canvas.paste(im,(x,y));draw.text((x+8,y+221),f'{i+1:02d}',fill='#142d3b')
    canvas.save(OUT/'contact-sheet.png')
manifest={'date':'2026-10-12','slides':18,'mainSlides':16,'backupSlides':2,'format':'16:9; native editable PowerPoint text/shapes; vector PDF','speakerNotes':18,'robotAccess':'none','modelAccess':'none','textOverflow':0,'sources':{str((ROOT/f'reports/ai-benchmark/{name}.json').relative_to(ROOT)):hashlib.sha256((ROOT/f'reports/ai-benchmark/{name}.json').read_bytes()).hexdigest() for name in results}}
(OUT/'validation.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'slides':18,'powerpointBytes':deck.stat().st_size,'pdfBytes':pdf_path.stat().st_size,'output':str(OUT)},ensure_ascii=False))

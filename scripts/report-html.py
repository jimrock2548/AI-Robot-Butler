"""Render report content to print HTML as an independent PDF fallback."""
import sys, html, json
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.tools/report-python'))
from docx import Document
from docx.oxml.ns import qn
OUT=ROOT/'reports/project-documentation'
document=Document(OUT/'AI-Robot-Butler-Report.docx')
page_map=json.loads((OUT/'toc-pages.json').read_text(encoding='utf-8')) if (OUT/'toc-pages.json').exists() else {}
headings=[p.text for p in document.paragraphs if p.style.name.startswith('Heading')]
def toc(names):
    return '<nav>'+''.join('<p style="text-indent:0;display:flex;justify-content:space-between"><span>'+html.escape(name)+'</span><span>'+str(page_map.get(name,'—'))+'</span></p>' for name in names)+'</nav>'
parts=[]
for element in document.element.body:
    if element.tag==qn('w:p'):
        text=''.join(element.itertext()) if False else ''.join(element.xpath('.//w:t/text()'))
        styles=element.xpath('./w:pPr/w:pStyle/@w:val');style=styles[0] if styles else 'Normal'
        if style.startswith('TOC') or element.xpath('.//w:instrText'):
            continue
        if element.xpath('.//w:br[@w:type="page"]'):
            parts.append('<div class="pagebreak"></div>')
        drawing=element.xpath('.//a:blip/@r:embed')
        if drawing:
            for rid in drawing:
                part=document.part.related_parts[rid]
                image_name=f'pdf-{rid}.png';(OUT/image_name).write_bytes(part.blob)
                parts.append(f'<img src="{image_name}" />')
        if not text:
            if not drawing and not element.xpath('.//w:br') and not element.xpath('./w:pPr/w:sectPr'):
                parts.append('<p class="empty">&nbsp;</p>')
            continue
        escaped=html.escape(text)
        if style.startswith('Heading'):
            level=style[-1] if style[-1].isdigit() else '2'
            cls=' class="chapter"' if text.startswith('บทที่ 1 ') else ''
            parts.append(f'<h{level}{cls}>{escaped}</h{level}>')
            if text=='สารบัญ':parts.append(toc([name for name in headings if name not in ['สารบัญ','สารบัญภาพ','สารบัญตาราง']]))
            elif text=='สารบัญภาพ':parts.append(toc([p.text for p in document.paragraphs if p.style.name=='Figure Caption']))
            elif text=='สารบัญตาราง':parts.append(toc([p.text for p in document.paragraphs if p.style.name=='Table Caption']))
        else:
            align=element.xpath('./w:pPr/w:jc/@w:val')
            css=' style="text-align:center;text-indent:0"' if align==['center'] else ''
            parts.append(f'<p{css}>{escaped}</p>')
    elif element.tag==qn('w:tbl'):
        rows=[]
        for row in element.xpath('./w:tr'):
            cells=[]
            for cell in row.xpath('./w:tc'):
                text='\n'.join(''.join(p.xpath('.//w:t/text()')) for p in cell.xpath('./w:p'))
                cells.append('<td>'+html.escape(text).replace('\n','<br/>')+'</td>')
            rows.append('<tr>'+''.join(cells)+'</tr>')
        parts.append('<table><thead>'+rows[0]+'</thead><tbody>'+''.join(rows[1:])+'</tbody></table>')
    elif element.tag==qn('w:sectPr'):pass
font_path=Path('C:/Users/admin/AppData/Local/Microsoft/Windows/Fonts/THSarabunNew.ttf').as_uri()
body='\n'.join(parts)
content=f'''<!doctype html><html lang="th"><meta charset="utf-8"><title>AI Robot Butler Report</title>
<style>@font-face{{font-family:Sarabun;src:url("{font_path}")}}
@page{{size:A4;margin:38.1mm 25.4mm 25.4mm 38.1mm;@bottom-center{{content:counter(page);font-family:Sarabun;font-size:14pt}}}}
body{{font-family:Sarabun,"TH Sarabun New",sans-serif;font-size:16pt;line-height:1.15;color:#111;margin:0;overflow-wrap:anywhere}}
h1{{font-size:20pt}}h2{{font-size:18pt}}h1,h2{{break-after:avoid}}.chapter{{break-before:page}}p{{margin:0 0 8pt;text-indent:1.25cm;break-inside:avoid}}.empty{{height:12pt;margin:0}}
table{{width:100%;table-layout:fixed;border-collapse:collapse;font-size:14pt;margin:10pt 0}}td{{border:1px solid #888;padding:4pt;vertical-align:top;overflow-wrap:anywhere}}thead{{display:table-header-group;font-weight:bold;background:#eee}}tr{{break-inside:avoid}}nav span:first-child{{min-width:0;flex:1}}nav span:last-child{{padding-left:12pt}}
img{{width:100%;break-inside:avoid;break-after:avoid}}.pagebreak{{break-before:page}}a{{color:inherit;text-decoration:none}}</style><body>{body}</body></html>'''
(OUT/'AI-Robot-Butler-Report.html').write_text(content,encoding='utf-8')
print('Print HTML ready')
